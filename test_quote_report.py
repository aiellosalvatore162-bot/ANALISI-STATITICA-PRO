"""Fase Z2: report quote partita × mercato dall'archivio unico (nessuna rete, nessun credito)."""

from datetime import date, datetime, timezone
import io

import pytest
from openpyxl import load_workbook

import day_archive
import quote_report
from match_memory import day_workbook, read_day_workbook
from quotes.model import QuoteMercato
from today import TodayMatch

T = "2026-10-30T10:20:00+00:00"


def _match(home, away, event):
    return TodayMatch(event, "soccer_italy_serie_a", "Serie A", datetime(2026, 10, 30, 18, tzinfo=timezone.utc), home, away)


def _q(partita, market, price, bookmaker="Eurobet", fonte="incollata", selection=None):
    return QuoteMercato(partita, market, None, selection or market, price, bookmaker, fonte, T)


@pytest.fixture()
def day():
    """Tre partite: Napoli-Roma completa, Inter-Milan con sospese, Lazio-Genoa senza quote."""
    store = {}
    keys = [day_archive.add_match(store, _match(h, a, e)) for h, a, e in
            (("Napoli", "Roma", "e1"), ("Inter", "Milan", "e2"), ("Lazio", "Genoa", "e3"))]
    for index, key in enumerate(keys):
        day_archive.set_selection(store, key, selected=True, to_play=index < 2)
    full = [_q("Napoli - Roma", key, 2.0 + index / 100) for index, (_, key) in enumerate(quote_report.columns())
            if key != quote_report.EXACT_KEY]
    full.append(_q("Napoli - Roma", "2-1", 9.0))
    day_archive.add_quotes(store, keys[0], full)
    day_archive.add_quotes(store, keys[1], [
        _q("Inter - Milan", "1", 1.9, "pinnacle", "The Odds API"), _q("Inter - Milan", "X", 3.4, "pinnacle", "The Odds API"),
        _q("Inter - Milan", "2", None), _q("Inter - Milan", "Goal", 1.8)])
    day_archive.add_page(store, keys[1], text="x", title="Inter - Milan", bookmaker="Eurobet", page_format="eurobet")
    return store, keys


def test_cells_ok_and_missing_with_reasons(day):
    store, keys = day
    rows = {row.partita: row for row in quote_report.build(store)}
    assert [row.partita for row in quote_report.build(store)] == ["Napoli - Roma", "Inter - Milan", "Lazio - Genoa"]  # Da giocare prima
    full = rows["Napoli - Roma"]
    assert full.complete and full.cells["1"].ok and full.cells["1"].text == "✓ 2.00 · Eurobet · incollata · 10:20"
    assert full.cells[quote_report.EXACT_KEY].text.startswith("✓ 1 punteggi")
    partial = rows["Inter - Milan"]
    assert partial.cells["1"].text == "✓ 1.90 · Pinnacle · The Odds API · 10:20"
    assert partial.cells["2"].motivo == "sospesa dal bookmaker" and partial.cells["2"].text.startswith("✗")
    assert partial.cells["Over 2.5 gol"].motivo == "assente nella pagina incollata"
    assert partial.cells["Multigol Casa 1-2"].motivo == "assente nella pagina incollata"
    empty = rows["Lazio - Genoa"]  # non «Da giocare», nessuna quota, nessuna pagina
    assert empty.cells["1"].motivo == "non raccolta (livello A)"
    assert empty.cells["Goal"].motivo.startswith("non raccolta: GG/NG solo per evento")
    assert empty.cells["Over 9.5 corner"].motivo == "fuori livello B (partita non «Da giocare»)"
    assert empty.cells["Multigol Casa 1-2"].motivo.startswith("mercato non offerto da The Odds API")
    assert not empty.complete and len(empty.missing) == len(quote_report.columns())


def test_failed_collection_reason_comes_from_the_ledger_rows(day):
    store, _ = day
    failed = [{"Partita": "Lazio - Genoa", "Livello": "A", "Mercato": "h2h, totals", "Motivo": "partita assente nella risposta"},
              {"Partita": "Lazio - Genoa", "Livello": "A", "Mercato": "GG/NG", "Motivo": "solo per evento (1 credito a partita)"}]
    row = quote_report.build(store, failed)[2]
    assert row.cells["1"].motivo == "non raccolta: partita assente nella risposta"
    assert row.cells["Goal"].motivo == "non raccolta: partita assente nella risposta"  # il motivo generico "solo per evento" si salta


def test_paste_list_orders_to_play_first_and_suggests_bookmakers(day):
    store, _ = day
    rows = quote_report.build(store)
    items = quote_report.to_paste(store, rows)
    assert [item["Partita"] for item in items][0] == "Inter - Milan" and items[-1]["Partita"] == "Lazio - Genoa"
    assert not any(item["Partita"] == "Napoli - Roma" for item in items)
    by = {(item["Partita"], item["Mercato"]): item for item in items}
    assert by[("Inter - Milan", "Over 2.5 gol")]["Bookmaker suggerito"] == "Eurobet"   # già offerto per Napoli - Roma
    assert by[("Inter - Milan", "2")]["Bookmaker suggerito"] == "Eurobet"
    text = quote_report.paste_text(items)
    assert text.startswith("Inter - Milan (Da giocare)") and "\nLazio - Genoa\n" in text
    assert "[suggerito: Eurobet]" in text


def test_summary_matches_ledger_and_projection(day):
    store, _ = day
    ledger = [{"Data": "2026-10-30", "Livello": "A", "Reali": 6}, {"Data": "2026-10-30", "Livello": "B", "Reali": 10},
              {"Data": "2026-10-30", "Livello": "B", "Reali": None},       # non riuscita: nessun credito
              {"Data": "2026-10-12", "Livello": "A", "Reali": 40}, {"Data": "2026-09-30", "Livello": "A", "Reali": 99}]
    info = quote_report.summary(store, quote_report.build(store), ledger, date(2026, 10, 30), budget=500)
    assert (info["crediti_oggi"], info["crediti_mese"], info["residui_mese"]) == (16, 56, 444)
    assert info["proiezione"]["totale_mese"] == 480 and info["proiezione"]["entro_budget"]
    assert (info["complete"], info["incomplete"]) == (1, 2)
    assert info["per_fonte"]["The Odds API"] == 2 and info["per_fonte"]["incollata"] == sum(
        1 for q in day_archive.quotes_of(store[next(iter(store))]) if q.quota is not None) + 1
    assert info["per_bookmaker"]["pinnacle"] == 2


def test_report_never_calls_credit_endpoints_and_follows_new_quotes(day, monkeypatch):
    import datasources.oddsapi as oddsapi

    def forbidden(*args, **kwargs):
        raise AssertionError("il report ha chiamato The Odds API")

    monkeypatch.setattr(oddsapi.OddsApi, "_get", forbidden)
    store, keys = day
    assert not quote_report.build(store)[2].cells["1"].ok
    day_archive.add_quotes(store, keys[2], [_q("Lazio - Genoa", "1", 2.2, "Snai")])   # un incolla: il report si aggiorna da solo
    cell = quote_report.build(store)[2].cells["1"]
    assert cell.ok and cell.bookmaker == "Snai" and "Snai" in cell.text
    quote_report.summary(store, quote_report.build(store))


def test_excel_sheet_exported_and_restored_without_differences(day):
    store, _ = day
    rows = quote_report.table(quote_report.build(store))
    content = day_workbook([], [], quotes_archive=store, quote_report=rows)
    sheet = load_workbook(io.BytesIO(content))["Report quote"]
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert values[0] == quote_report.sheet_columns()
    assert [dict(zip(values[0], row)) for row in values[1:]] == [{column: row.get(column) or None for column in values[0]} for row in rows]
    restored = {}
    day_archive.restore(restored, read_day_workbook(content).archive)
    assert quote_report.table(quote_report.build(restored)) == rows


UI_SCRIPT = '''
import sys
sys.path.insert(0, {root!r})
import streamlit as st
import day_archive, ui_quotes
from datetime import datetime, timezone
from quotes.model import QuoteMercato
from today import TodayMatch
class Api:
    settings = {{"budget_mensile": 500}}
    class cache:
        @staticmethod
        def _connect():
            raise RuntimeError("niente registro")
store = ui_quotes.archive()
key = day_archive.add_match(store, TodayMatch("e", "soccer_italy_serie_a", "Serie A", datetime(2026, 10, 30, 18, tzinfo=timezone.utc), "Napoli", "Roma"))
day_archive.set_selection(store, key, selected=True)
day_archive.add_quotes(store, key, [QuoteMercato("Napoli - Roma", "1", None, "1", 2.1, "Eurobet", "incollata", "2026-10-30T10:20:00+00:00")])
ui_quotes.render_quote_report(Api())
'''


def test_report_expander_renders_in_the_app(tmp_path):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest

    script = tmp_path / "report.py"
    script.write_text(UI_SCRIPT.format(root=str(Path(__file__).parent)))
    at = AppTest.from_file(str(script), default_timeout=60).run()
    assert not at.exception
    assert at.expander[0].label == "Report quote · 0 complete, 1 incomplete"
    assert any("Da incollare a mano" in markdown.value for markdown in at.markdown)
