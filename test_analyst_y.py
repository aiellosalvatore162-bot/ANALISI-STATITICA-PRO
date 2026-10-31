"""Fase Y: report completo dell'Analista (ESITO, PARTITA, PROFILO_TATTICO), storico proprio, regole dell'Analista,
forza per il calendario senza quote e confronto Eurobet con lo sharp. Solo dati locali, nessuna rete."""

from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

import analyst
import analyst_dataset
import engine
import sharp_value
import tactical_sheets
from datasources.cache import SourceCache
from engine import config_from_input, simulate_match
from match_memory import day_workbook, read_day_workbook
from parser import parse_match_input
from performance import settle
from prediction_log import COLUMNS, SaveInfo, prediction_rows
from slips import PENDING, WON, Slip
from tactical_sheets import TacticalSheet
from test_engine import VALID

REPORT = (Path(__file__).parent / "tests" / "fixtures" / "analista" / "report_completo.md").read_text(encoding="utf-8")


def _row(**values):
    row = {column: None for column in COLUMNS}
    row.update({"Ora": "18:00", "Tipo": "A carta", "Versione_modello": "test", "Note": "", "Prob_finale": 0.5,
                "Quota_presa": 2.0}, **values)
    return row


def test_full_report_reads_every_block():
    report = analyst.parse_full(REPORT, date(2026, 10, 26))
    assert [block.label for block in report.matches] == ["Genoa - Venezia", "Real Madrid - Barcellona"]
    genoa = report.matches[0]
    # Il blocco PARTITA finisce dove inizia il successivo: niente profili o esiti dentro le correzioni.
    assert genoa.response.absent == {"ospite": ["Mario Rossi"]} and genoa.keys()[0]
    assert "PROFILO_TATTICO" not in genoa.corrections_text
    outcomes = {block.label: block for block in report.outcomes}
    assert set(outcomes) == {"Inter - Fiorentina", "AS Roma - Napoli", "Real Madrid - Juventus", "Lecce - Juventus"}
    inter = outcomes["Inter - Fiorentina"]
    assert inter.values["away_double_yellows"] == 1 and inter.values["home_corners_1t"] == 4 and inter.has_halves
    assert outcomes["Real Madrid - Juventus"].competition == "Champions League"
    assert outcomes["Lecce - Juventus"].error.startswith("senza fonte")
    assert [sheet.team for sheet in report.profiles] == ["Genoa", "Venezia"]
    assert all(sheet.origine == "Analista" and sheet.aggiornata == "2026-10-26" for sheet in report.profiles)
    assert any("Barcellona" in warning for warning in report.warnings)


def test_outcomes_fill_registry_and_slips_without_overwriting():
    outcomes = analyst.parse_outcomes(REPORT)
    rows = [_row(ID="a", Data="2026-10-25", Partita="Inter - Fiorentina", Mercato="Corner per tempo", Selezione="Over 4.5 corner 1T"),
            _row(ID="b", Data="2026-10-25", Partita="Inter - Fiorentina", Mercato="Cartellini per tempo", Selezione="Over 1.5 cartellini 1T"),
            # Punti Eurobet: gialli 2 + 3 (con il primo giallo della doppia ammonizione) + rosso 2 = 7.
            _row(ID="c", Data="2026-10-25", Partita="Inter - Fiorentina", Mercato="Cartellini", Selezione="Over 6.5 cartellini"),
            _row(ID="d", Data="2026-10-22", Partita="Real Madrid - Juventus", Mercato="Corner", Selezione="Over 9.5 corner"),
            _row(ID="e", Data="2026-10-25", Partita="AS Roma - Napoli", Mercato="1X2", Selezione="X", Esito=0.0, Note="a mano"),
            _row(ID="f", Data="2026-10-25", Partita="Lecce - Juventus", Mercato="1X2", Selezione="X")]
    changed, settled = analyst.fill_registry(rows, outcomes, today=date(2026, 10, 26))
    by_id = {row["ID"]: row for row in changed}
    assert settled == 4 and set(by_id) == {"a", "b", "c", "d"}
    assert (by_id["a"]["Esito"], by_id["b"]["Esito"], by_id["c"]["Esito"], by_id["d"]["Esito"]) == (1.0, 0.0, 1.0, 1.0)
    assert "Analista (https://www.legaseriea.it/it/match/inter-fiorentina)" in by_id["a"]["Note"]
    slip = Slip("Weekend", "2026-10-24T10:00", [{"match": "AS Roma - Napoli", "date": "2026-10-25", "selection": "Under 2.5 gol",
                                                  "odds": 1.8, "probability": 0.5, "status": PENDING, "note": ""}])
    slips, legs = analyst.fill_slips([slip.to_dict()], outcomes)
    assert legs == 1 and slips[0]["legs"][0]["status"] == WON


def test_double_yellows_are_used_when_the_bookmaker_rule_needs_them():
    rule = engine.CardRule(double_yellow="due_gialli_e_rosso", double_yellow_share=0.5)
    result = {"home_goals": 1, "away_goals": 0, "home_yellows": 2, "away_yellows": 3, "home_reds": 0, "away_reds": 1}
    assert settle("Over 6.5 cartellini", result, rule) is None
    # Con la doppia ammonizione indicata: 2 + 3 + 2 (rosso) + 1 (secondo giallo) = 8 punti.
    assert settle("Over 6.5 cartellini", {**result, "home_double_yellows": 0, "away_double_yellows": 1}, rule) == 1


def test_profiles_ask_for_confirmation_only_when_a_value_changes():
    report = analyst.parse_full(REPORT, date(2026, 10, 26))
    saved = [TacticalSheet("Venezia", blocco="basso", note="a mano", fonti=["https://a.it"], origine="manuale").to_dict(),
             TacticalSheet("Genoa", modulo="4-2-3-1", fonti=["https://b.it"]).to_dict()]
    sheets, changes, pending = analyst.merge_profiles(saved, report.profiles)
    assert changes == [{"Squadra": "Venezia", "Campo": "blocco", "Prima": "basso", "Dopo": "alto"}]
    assert [sheet.team for sheet in pending] == ["Venezia"] and pending[0].note == "a mano"
    genoa = next(item for item in sheets if item["team"] == "Genoa")
    assert genoa["transizione"] == 4 and genoa["origine"] == "Analista" and "https://b.it" in genoa["fonti"]
    venezia = next(item for item in sheets if item["team"] == "Venezia")
    assert venezia["blocco"] == "basso"  # cambia solo dopo la conferma


def test_analyst_rules_are_small_flagged_and_need_profiles():
    report = analyst.parse_full(REPORT, date(2026, 10, 26))
    genoa, venezia = report.profiles
    factors, notes = tactical_sheets.analyst_adjustment(genoa, venezia)
    # Fasce/cross alto contro difesa a uomo: 1.08 attenuato (adattabilità media) e tagliato a 1.05.
    assert factors["corner"] == pytest.approx(1.05) and factors["gol_casa"] == 1.0
    assert notes and "non ancora verificata" in notes[0]
    manual = replace(genoa, origine="manuale")
    assert tactical_sheets.analyst_adjustment(manual, venezia)[1] == []
    effects = tactical_sheets.match_effects([genoa.to_dict(), venezia.to_dict()], "Genoa", "Venezia")
    assert effects.analyst_rules and effects.factors["corner"] == pytest.approx(1.05)
    off = tactical_sheets.load_rules()
    off["regole_analista"]["attiva"] = False
    assert tactical_sheets.analyst_adjustment(genoa, venezia, off)[1] == []


def test_registry_notes_for_analyst_rules_and_eurobet_vs_sharp():
    result = simulate_match(config_from_input(parse_match_input(VALID), simulations=2000, seed=3))
    result["analyst_rules"] = True
    info = SaveInfo(date(2026, 10, 26), "15:00", "Serie A", "A carta", "test")
    row = prediction_rows(result, ["1"], info, extra_notes={"1": sharp_value.note(4.2)})[0]
    assert row["Note"] == "motore: Calcolo attuale; regole_analista: sì; eurobet_vs_sharp: +4.2%"


def test_eurobet_against_the_fair_sharp_price():
    sharp = {"1": 2.0, "X": 3.6, "2": 4.0, "Over 2.5 gol": 1.9, "Under 2.5 gol": 1.95}
    eurobet = {"1": 2.15, "X": 3.3, "2": 4.1, "1X": 1.35, "Casa DNB": 1.5, "Over 2.5 gol": 2.0}
    rows = {row["Selezione"]: row for row in sharp_value.compare(eurobet, sharp)}
    fair = sharp_value.fair_probabilities(sharp)
    assert sum(fair[key] for key in ("1", "X", "2")) == pytest.approx(1.0)
    assert rows["1"]["Scarto %"] == pytest.approx((2.15 * fair["1"] - 1) * 100, abs=0.1)
    assert rows["X"]["Scarto %"] < 0 and "sotto" in sharp_value.describe(rows["X"])
    assert set(rows) == {"1", "X", "2", "1X", "Casa DNB", "Over 2.5 gol"}
    above = sharp_value.above(list(rows.values()), 3.0)
    assert [row["Selezione"] for row in above] == sorted((row["Selezione"] for row in above),
                                                         key=lambda key: -rows[key]["Scarto %"])
    assert all(row["Scarto %"] >= 3.0 for row in above)
    assert sharp_value.compare(eurobet, None) == [] and sharp_value.threshold_pct() == 3.0
    frame = pd.DataFrame({"Note": ["eurobet_vs_sharp: +5.0%", "eurobet_vs_sharp: -2.0%", ""], "Esito": [1.0, 0.0, 1.0],
                          "Quota_presa": [2.2, 2.0, 1.5], "Quota_chiusura": [2.0, 2.1, None]})
    table = sharp_value.performance_table(frame).set_index("Gruppo")
    assert table.iloc[0]["Selezioni"] == 1 and table.iloc[0]["Rendimento medio"] == pytest.approx(1.2)
    assert table.iloc[0]["CLV medio"] == pytest.approx(0.1)


def test_schedule_strength_only_without_market_odds():
    config = config_from_input(parse_match_input(VALID), simulations=2000, seed=3)
    strengths = {"home": (1.6, 1.4), "away": (0.7, 0.8)}
    v2 = replace(config.v2, strength_model_only=True, strength_goal_params=config.goal_params)
    no_odds = replace(config, odds={}, anchor_odds=None, strength_inputs=strengths, v2=v2)
    lam, mu, details = engine.final_goal_rates(no_odds)
    assert "forza_calendario" in details["ledger"] and details["ledger"]["forza_calendario"]["ingressi"] == strengths
    plain = engine.final_goal_rates(replace(no_odds, strength_inputs=None))
    assert lam > plain[0] and mu < plain[1]
    with_odds = replace(no_odds, odds={"1": 2.0, "X": 3.4, "2": 3.8})
    assert "forza_calendario" not in engine.final_goal_rates(with_odds)[2]["ledger"]


def _record(index: int, applies: bool, goals: int, corners_1t: int = 4):
    home = TacticalSheet(f"Casa{index}", transizione=5 if applies else 1, fasce_cross="alto" if applies else "basso",
                         origine="Analista")
    away = TacticalSheet(f"Ospite{index}", blocco="alto", non_possesso="uomo" if applies else "zona", origine="Analista")
    return {"id": f"r{index}", "data": f"2026-{1 + index // 28:02d}-{1 + index % 28:02d}", "competizione": "Serie A",
            "casa": home.team, "ospite": away.team, "profilo_casa": home.to_dict(), "profilo_ospite": away.to_dict(),
            "quote": {}, "probabilita": {"Over 4.5 corner 1T": 45.0, "Over 1.5 cartellini 1T": 40.0},
            "attesi": {"lambda": 1.5, "mu": 1.1, "corner": 9.5, "gialli": 4.2}, "regole": {},
            "esito": {"home_goals": goals, "away_goals": 1, "home_corners": 6, "away_corners": 4, "home_yellows": 2,
                      "away_yellows": 2, "home_reds": 0, "away_reds": 0, "home_corners_1t": corners_1t,
                      "away_corners_1t": 1, "home_yellows_1t": 1, "away_yellows_1t": 0, "home_reds_1t": 0, "away_reds_1t": 0}}


def test_dataset_keeps_first_profiles_and_never_overwrites_outcomes(tmp_path):
    result = simulate_match(config_from_input(parse_match_input(VALID), simulations=2000, seed=3))
    first = analyst_dataset.from_result(date(2026, 10, 25), "Serie A", "Inter", "Fiorentina", result, {"1": 1.6},
                                        TacticalSheet("Inter", blocco="alto"), None)
    assert first["attesi"]["lambda"] == result["model"]["lambda"] and "1" in first["probabilita"]
    assert all(key in analyst_dataset.MAIN_KEYS or key.endswith((" 1T", " 2T")) for key in first["probabilita"])
    later = analyst_dataset.from_result(date(2026, 10, 25), "Serie A", "Inter", "Fiorentina", result, {"1": 1.7},
                                        TacticalSheet("Inter", blocco="basso"), TacticalSheet("Fiorentina"))
    records = analyst_dataset.upsert(analyst_dataset.upsert([], first), later)
    assert len(records) == 1 and records[0]["profilo_casa"]["blocco"] == "alto" and records[0]["profilo_ospite"]
    outcome = next(block for block in analyst.parse_outcomes(REPORT) if block.label == "Inter - Fiorentina")
    records, status = analyst_dataset.attach_outcome(records, outcome)
    assert status == "aggiunto" and records[0]["esito"]["home_corners_1t"] == 4
    assert analyst_dataset.attach_outcome(records, outcome)[1] == "già presente"
    cup = next(block for block in analyst.parse_outcomes(REPORT) if block.label == "Real Madrid - Juventus")
    records, status = analyst_dataset.attach_outcome(records, cup)
    assert status.startswith("nuova partita") and len(records) == 2
    store = analyst_dataset.DatasetStore(SourceCache(tmp_path / "cache.sqlite", sleep=lambda _: None, min_interval={}))
    store.save(records)
    assert store.load() == sorted(records, key=lambda item: (item["data"], item["id"]))
    restored = read_day_workbook(day_workbook([], [], dataset=records))
    assert [item["id"] for item in restored.dataset] == [item["id"] for item in records]
    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(day_workbook([], [], dataset=records)))["Dataset Analista"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == analyst_dataset.SHEET_COLUMNS and rows[1][7] == "2-1" and rows[1][10] == "4-1"


def test_rule_verification_waits_for_enough_matches_and_then_judges():
    few = [_record(index, index % 2 == 0, 1) for index in range(20)]
    assert analyst_dataset.verify_rules(few)["regole"] == {}
    rng = np.random.default_rng(4)
    many = [_record(index, index % 2 == 0, int(rng.poisson(1.5)), int(rng.poisson(4))) for index in range(200)]
    result = analyst_dataset.verify_rules(many)
    rule = result["regole"]["fasce_vs_difesa_uomo"]
    assert result["partite"] == 200 and rule["casi"] == 100 and "corner" in rule["verifica"]
    assert 0.95 <= rule["effetto"]["corner"] <= 1.05
    assert rule["proposta"] in ("passa a «verificata»", "spegni", "lascia «da verificare»")


def test_rule_proposal_is_written_only_on_confirmation(tmp_path):
    path = tmp_path / "tattica.json"
    path.write_text(tactical_sheets.CONFIG_PATH.read_text(encoding="utf-8"), encoding="utf-8")
    verdict = {"proposta": "spegni", "partite": 160, "casi": 40, "effetto": {"corner": 0.97},
               "verifica": {"corner": {"differenza": 0.01, "due_es": 0.005, "ece_prima": 0.02, "ece_dopo": 0.03,
                                       "migliora": False, "peggiora": True}}}
    analyst_dataset.apply_proposal("fasce_vs_difesa_uomo", verdict, path)
    settings = tactical_sheets.load_rules(path)
    assert settings["regole_analista"]["regole"]["fasce_vs_difesa_uomo"]["stato"] == "spenta"
    genoa, venezia = analyst.parse_full(REPORT).profiles
    assert tactical_sheets.analyst_adjustment(genoa, venezia, settings)[1] == []


def test_half_markets_are_calibrated_once_the_dataset_is_large_enough():
    records = [_record(index, False, 1, corners_1t=4 if index % 2 else 2) for index in range(120)]
    assert analyst_dataset.half_calibration(records[:50], 100)["mercati"] == []
    table = {row["Mercato"]: row for row in analyst_dataset.half_calibration(records, 100)["mercati"]}
    # Corner 1T: 4+1 = 5 > 4.5 nella metà delle partite; cartellini 1T: 1 punto, mai sopra 1.5.
    assert table["Over 4.5 corner 1T"]["Reale %"] == 50.0 and table["Over 4.5 corner 1T"]["Prevista %"] == 45.0
    assert table["Over 1.5 cartellini 1T"]["Reale %"] == 0.0
