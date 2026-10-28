"""Memoria delle partite analizzate ed Excel della giornata (fase K): nessuna rete, simulazioni piccole."""

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
import pytest

import engine
from datasources.cache import SourceCache
from tactical_sheets import SHEET_COLUMNS, TacticalSheet
from match_memory import (
    DUAL_SHEET, REPORT_SHEET, SLIPS_SHEET, ANALYSIS_SHEET, REGISTRY_SHEET, TACTICAL_SHEET, TECHNICAL_SHEET, analysis_rows, day_workbook, make_record, match_id,
    read_day_workbook, summary_label, upsert,
)
from prediction_log import COLUMNS, PredictionLog, SaveInfo, prediction_rows


TEMPLATE = Path(__file__).parent / "tests" / "fixtures" / "Registro_Previsioni.xlsx"
EUROBET = Path(__file__).parent / "tests" / "fixtures" / "eurobet" / "eurobet_corner.txt"
MATCHES = (("Napoli", "Roma", 1850, 1800, {"1": 2.10, "X": 3.30, "2": 3.60, "Over 2.5 gol": 1.95, "Under 2.5 gol": 1.85}),
           ("Inter", "Milan", 1900, 1840, {"1": 1.95, "X": 3.50, "2": 3.90}),
           ("Lazio", "Torino", 1780, 1700, {"1": 1.80, "X": 3.40, "2": 4.80, "Over 9.5 corner": 2.10}))
DAY = date(2026, 10, 25)


def block(home, away):
    team = "Attacco: 1.2 | Difesa: 1.1 | Elo: 1850 | Forma: W-W-D-W-L | xG Fatti: 1.8 | xG Subiti: 1.0 | " \
           "Corner pro/sub: 5.5/4.0 | Cartellini pro: 2.0"
    return f"CASA: {home}\n{team}\nOSPITE: {away}\n{team}\nARBITRO: Gialli medi 4.2 | Falli medi 24\n"


def simulate(home, away, home_elo, away_elo, odds):
    team = lambda name, elo: engine.TeamParams(name, 1.1, 1.0, elo, "W-D-W-L-W", 1.5, 1.2, 5.5, 4.5, 2.0)
    config = engine.MatchConfig(team(home, home_elo), team(away, away_elo), 4.4, 24.0, odds=odds, simulations=4_000)
    return engine.simulate_match(config)


@pytest.fixture(scope="module")
def day():
    """Tre partite simulate in sequenza: memoria e righe del registro della giornata."""
    store, registry = {}, []
    info = SaveInfo(DAY, "20:45", "Serie A", "A carta", "test")
    for home, away, home_elo, away_elo, odds in MATCHES:
        result = simulate(home, away, home_elo, away_elo, odds)
        eurobet = EUROBET.read_text(encoding="utf-8") if home == "Lazio" else ""
        record = make_record(result, block=block(home, away), odds=odds, day=DAY, competition="Serie A",
                             eurobet_text=eurobet, warnings=["prova"], claude_answer="", simulations=4_000, seed=42,
                             saved_at=datetime(2026, 10, 25, 12, 0))
        assert upsert(store, record) is False
        registry += prediction_rows(result, ["1", "Over 2.5 gol"], info, odds=odds)
    registry.append({**{column: None for column in COLUMNS}, "ID": "2026-10-25|Multipla|Multipla|1 + 1 + Over 2.5",
                     "Data": DAY.isoformat(), "Partita": "Multipla 3 partite", "Mercato": "Multipla",
                     "Selezione": "Napoli 1 + Inter 1 + Lazio Over 2.5", "Prob_finale": 0.18, "Quota_presa": 6.4,
                     "Tipo": "A carta"})
    return store, registry


def test_three_matches_in_sequence_give_three_matches_in_the_file(day):
    store, registry = day
    content = day_workbook(registry, store.values())
    workbook = load_workbook(__import__("io").BytesIO(content))
    assert workbook.sheetnames == [REGISTRY_SHEET, ANALYSIS_SHEET, DUAL_SHEET, REPORT_SHEET, SLIPS_SHEET, TACTICAL_SHEET,
                                  "Dataset Analista", TECHNICAL_SHEET]
    assert workbook[TECHNICAL_SHEET].sheet_state == "hidden"
    analysis = list(workbook[ANALYSIS_SHEET].iter_rows(values_only=True))
    assert len(analysis) == 1 + 3
    assert [row[3] for row in analysis[1:]] == ["Napoli - Roma", "Inter - Milan", "Lazio - Torino"]
    header = analysis[0]
    assert {"P 1", "Quota 1", "EV 1", "P Over 2.5 gol", "P Goal", "P Multigol 2-4", "P 1 1T", "P Over 9.5 corner",
            "P Over 4.5 cartellini", "Corner attesi", "Cartellini attesi"} <= set(header)
    napoli = dict(zip(header, analysis[1]))
    assert napoli["Quota 1"] == 2.10 and 0 < napoli["P 1"] < 1 and napoli["EV 1"] is not None
    registry_sheet = list(workbook[REGISTRY_SHEET].iter_rows(values_only=True))
    assert tuple(registry_sheet[0]) == COLUMNS and len(registry_sheet) == 1 + 7
    assert "Multipla" in {row[5] for row in registry_sheet[1:]}
    assert list(workbook[TACTICAL_SHEET].iter_rows(values_only=True)) == [SHEET_COLUMNS]


def test_resume_from_file_restores_registry_and_reopenable_matches(day):
    store, registry = day
    sheet = TacticalSheet("Napoli", transizione=4, blocco="alto", modulo="4-3-3", dal="2026-07-01",
                          fonti=["https://esempio.it/napoli"], origine="Claude, confermata")
    content = day_workbook(registry, store.values(), [sheet.to_dict()])
    restored = read_day_workbook(content)
    visible = list(load_workbook(__import__("io").BytesIO(content))[TACTICAL_SHEET].iter_rows(values_only=True))
    row = dict(zip(visible[0], visible[1]))
    assert row["Squadra"] == "Napoli" and row["Transizione"] == 4 and row["Fonti"] == "https://esempio.it/napoli"
    assert [record["id"] for record in restored.records] == list(store)
    lazio = restored.records[2]
    assert lazio == store[match_id(DAY, "Lazio", "Torino")]
    assert lazio["block"].startswith("CASA: Lazio") and lazio["eurobet_text"].startswith("1X2 CORNER")
    assert lazio["odds"]["Over 9.5 corner"] == 2.10
    assert len(restored.registry_rows) == 7 and TacticalSheet.from_dict(restored.tactical_rows[0]) == sheet
    multipla = next(row for row in restored.registry_rows if row["Mercato"] == "Multipla")
    assert multipla["Quota_presa"] == 6.4 and multipla["Prob_finale"] == 0.18


def test_reanalysing_the_same_match_updates_the_entry(day):
    store, _ = day
    store = dict(store)
    again = make_record(simulate("Napoli", "Roma", 1850, 1800, {"1": 2.00}), block=block("Napoli", "Roma"),
                        odds={"1": 2.00}, day=DAY, competition="Serie A")
    assert upsert(store, again) is True and len(store) == 3
    assert store[match_id(DAY, "Napoli", "Roma")]["odds"] == {"1": 2.00}


def test_session_registry_accumulates_across_matches_and_deletes_rows(day, tmp_path):
    _, registry = day
    log = PredictionLog(SourceCache(tmp_path / "cache.sqlite"))
    log.save(registry[:2], "s1")       # prima partita
    log.save(registry[2:6], "s1")      # altre due: non si azzera
    assert summary_label(log.rows("s1")) == "6 selezioni da 3 partite"
    assert log.save(registry[:2], "s1") == (0, 2)  # stesso ID: nessun doppione
    assert log.delete([registry[0]["ID"]], "s1") == 1 and len(log.rows("s1")) == 5
    assert log.delete([registry[1]["ID"]], "altra-sessione") == 0


def test_long_texts_are_split_across_cells(day):
    store, registry = day
    record = dict(next(iter(store.values())), eurobet_text="x" * 70_000)
    restored = read_day_workbook(day_workbook(registry, [record]))
    assert restored.records[0]["eurobet_text"] == "x" * 70_000


def test_user_registry_file_without_technical_sheet():
    restored = read_day_workbook(TEMPLATE.read_bytes())
    assert restored.records == [] and restored.registry_rows == []
    assert summary_label([]) == "0 selezioni da 0 partite"


def test_analysis_rows_flag_claude_and_eurobet(day):
    store, _ = day
    rows = analysis_rows(store.values())
    assert [row["Eurobet"] for row in rows] == ["", "", "incollata"] and rows[0]["Segnalazioni"].startswith("prova")
