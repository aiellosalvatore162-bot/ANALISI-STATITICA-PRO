"""Report dell'Analista del giorno (fase W): lettura, abbinamento alle partite di Oggi, regole della fase G."""

from datetime import date
from pathlib import Path

import pytest

import analyst


REPORT = (Path(__file__).parent / "tests" / "fixtures" / "analista" / "report_esempio.md").read_text(encoding="utf-8")
TARGETS = [analyst.Target(0, "Inter Milan", "Fiorentina", date(2026, 10, 25)),
           analyst.Target(1, "AS Roma", "Napoli", date(2026, 10, 25)),
           analyst.Target(2, "Genoa", "Venezia", date(2026, 10, 25)),
           analyst.Target(3, "Juventus", "Lecce", date(2026, 10, 25))]


@pytest.fixture
def blocks():
    return analyst.match_blocks(analyst.parse_report(REPORT), TARGETS)


def test_all_blocks_are_found_among_free_text(blocks):
    assert [block.label for block in blocks] == ["Inter - Fiorentina", "AS Roma - Napoli", "Genoa - Venezia",
                                                 "Real Madrid - Barcellona", "Lecce - Juventus", "Inter Milan - Fiorentina"]
    first = blocks[0]
    assert first.competition == "Serie A" and first.when.isoformat() == "2026-10-25T18:00:00"
    assert first.comment.startswith("L'Inter ruota") and "Parentesi" not in first.corrections_text
    assert first.corrections_text.startswith("CORREZIONI_CLAUDE:") and first.corrections_text.endswith("Under per la Fiorentina.")


def test_matching_by_team_names_written_differently_and_date(blocks):
    # Genoa - Venezia: report del 26/10, pagina del 25/10, dentro il giorno di tolleranza per il fuso orario.
    assert [block.matched for block in blocks] == [0, 1, 2, None, None, None]
    real, lecce, duplicate = blocks[3:]
    assert real.reason == "nessuna partita della pagina Oggi con queste squadre"
    assert lecce.reason == "manca il blocco CORREZIONI_CLAUDE"
    assert duplicate.reason == "partita già abbinata al blocco «PARTITA: Inter - Fiorentina | Serie A | 25/10/2026 18:00»: uso solo il primo"


def test_date_tolerance_and_reversed_teams():
    blocks = analyst.parse_report("PARTITA: Genoa - Venezia | Serie A | 28/10/2026 15:00\nCORREZIONI_CLAUDE:\nconfidenza: bassa\n"
                                  "PARTITA: Juventus - Lecce | Serie A | 25/10/2026\nCORREZIONI_CLAUDE:\nconfidenza: bassa\n"
                                  "PARTITA: Fiorentina - Inter | Serie A | 25/10/2026\nCORREZIONI_CLAUDE:\nconfidenza: bassa\n")
    genoa, juve, reversed_ = analyst.match_blocks(blocks, TARGETS)
    assert genoa.matched is None and genoa.reason == "data diversa: report 28/10/2026, pagina 25/10/2026"
    assert juve.matched == 3
    assert reversed_.matched is None and "ordine inverso" in reversed_.reason


def test_phase_g_rules_bounds_sources_and_referee(blocks):
    inter, roma = blocks[0], blocks[1]
    assert inter.response.adjustment().goals_away == 0.93
    referee = inter.response.referee_yellow(4.5)
    assert referee == pytest.approx((6 * 5.4 + 10 * 4.5) / 16)          # avvicinata alla media con peso 10 partite
    adjustment = roma.response.adjustment()
    assert adjustment.goals_away == 1.10 and adjustment.cards == 1.10 and adjustment.corners == 1.0
    assert roma.applied() == ["gol_ospite 1.1", "cartellini 1.1"]
    assert roma.discarded() == ["corner (senza fonte: ignorata)"]
    assert roma.response.confidence == "alta"


def test_report_table(blocks):
    rows = analyst.report_rows(blocks, {0: "Inter Milan - Fiorentina", 1: "AS Roma - Napoli"})
    assert rows[1]["Partita abbinata"] == "AS Roma - Napoli" and rows[1]["Correzioni scartate"].startswith("corner")
    assert rows[3]["Partita abbinata"] == "" and rows[3]["Non abbinato perché"]


def test_report_is_saved_in_the_day_file(blocks):
    from io import BytesIO

    from openpyxl import load_workbook

    from match_memory import REPORT_SHEET, day_workbook, read_day_workbook

    rows = analyst.report_rows(blocks, {0: "Inter Milan - Fiorentina", 1: "AS Roma - Napoli", 2: "Genoa - Venezia"})
    report = {"id": "abc", "date": "2026-10-25", "saved_at": "2026-10-25T10:00:00", "text": REPORT, "rows": rows}
    content = day_workbook([], [], reports=[report])
    sheet = list(load_workbook(BytesIO(content))[REPORT_SHEET].iter_rows(values_only=True))
    assert len(sheet) == 1 + 6 and sheet[2][2] == "AS Roma - Napoli" and sheet[2][4] == "AS Roma - Napoli"
    assert sheet[4][8] == "nessuna partita della pagina Oggi con queste squadre"
    assert read_day_workbook(content).reports == [report]
