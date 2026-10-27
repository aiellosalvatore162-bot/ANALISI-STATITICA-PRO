"""Fase X: ricerca rapida, frequenza storica recente, schedine salvate e marcatori (solo fixture, nessuna rete)."""

from datetime import date, datetime
import json
from pathlib import Path

import pandas as pd
import pytest

import engine
import quick_search
import recent_form
import scorers
import slips
from multiples import Leg
from tools.dataset import clean_frames, frame_from_csv


SOURCES = Path(__file__).parent / "tests" / "fixtures" / "sources"


@pytest.mark.parametrize("text, market, family, line, count, threshold, band", [
    ("dammi 6 over 2.5 per stasera", "over", "gol", 2.5, 6, None, "stasera"),
    ("4 goal più probabili", "goal", "gol", None, 4, None, None),
    ("3 multigol casa sopra il 70%", "multigol casa", "gol", None, 3, 0.70, None),
    ("5 over 8.5 corner", "over", "corner", 8.5, 5, None, None),
    ("doppia chance sicure oggi", "doppia chance", "gol", None, quick_search.DEFAULT_COUNT, 0.75, "oggi"),
    ("sei under 3,5 nel pomeriggio", "under", "gol", 3.5, 6, None, "pomeriggio"),
    ("3 over cartellini almeno 65%", "over", "cartellini", 4.5, 3, 0.65, None),
    ("2 no goal", "no goal", "gol", None, 2, None, None),
])
def test_phrases_are_understood(text, market, family, line, count, threshold, band):
    query = quick_search.parse_query(text)
    assert (query.market, query.family, query.line, query.count, query.threshold, query.band) == \
        (market, family, line, count, threshold, band)


def test_unknown_phrases_return_none_and_examples_are_valid():
    assert quick_search.parse_query("che tempo fa domani") is None
    assert all(quick_search.parse_query(text) for text in quick_search.EXAMPLES + quick_search.PRESETS)
    assert quick_search.parse_query("multigol 1-3").low_high == (1, 3)


def simulated(home_elo, away_elo, odds=None):
    team = lambda name, elo: engine.TeamParams(name, 1.1, 1.0, elo, "W-D-W-L-W", 1.5, 1.2, 5.5, 4.5, 2.0)
    return engine.simulate_match(engine.MatchConfig(team("A", home_elo), team("B", away_elo), 4.4, 24.0, odds=odds,
                                                    simulations=4_000))


@pytest.fixture(scope="module")
def matches():
    rows = []
    for index, (home_elo, away_elo, hour) in enumerate([(1900, 1600, 20), (1750, 1750, 15), (1650, 1950, 21), (1800, 1700, 18)]):
        result = simulated(home_elo, away_elo)
        final = result["components"]["finale"]
        odds = {"Over 2.5 gol": 2.0, "1X": round(100 / final["1X"] * 0.97, 2), "X2": 1.01, "12": round(100 / final["12"], 2)}
        rows.append({"id": str(index), "label": f"Casa{index} - Ospite{index}", "kickoff": datetime(2026, 10, 25, hour, 0),
                     "result": result, "odds": odds})
    return rows


def test_build_picks_most_probable_selection_with_band_threshold_and_traps(matches):
    proposal = quick_search.build(quick_search.parse_query("3 doppia chance per stasera"), matches)
    assert len(proposal.legs) == 3 and all(match["kickoff"].hour >= 18 for match in matches if match["id"] in
                                           {leg.match_id for leg in proposal.legs})
    assert any("fuori dalla fascia" in note for note in proposal.notes)
    assert [leg.probability for leg in proposal.legs] == sorted((leg.probability for leg in proposal.legs), reverse=True)
    # X2 a 1.01 è una trappola: mai proposta.
    assert all(leg.selections != ("X2",) or leg.ev * 100 >= -5 for leg in proposal.legs)
    assert proposal.traps and all(leg.selections == ("X2",) for leg in proposal.traps)
    strict = quick_search.build(quick_search.parse_query("4 over 2.5 sopra il 99%"), matches)
    assert strict.legs == [] and any("Trovate solo 0" in note for note in strict.notes)
    multigol = quick_search.build(quick_search.parse_query("2 multigol casa"), matches)
    assert all(leg.selections[0].startswith("Multigol Casa") and leg.odds is None for leg in multigol.legs)


@pytest.fixture(scope="module")
def serie_a():
    body = (SOURCES / "footballdata_I1_2627.csv").read_bytes()
    return clean_frames([frame_from_csv(body, "I1", "2627")])


def test_recent_frequency_counts_last_matches_of_both_teams(serie_a):
    last_day = serie_a["date"].max().date()
    freq = recent_form.recent_frequency("Over 2.5 gol", "Inter Milan", "Napoli", serie_a, as_of=last_day + pd.Timedelta(days=1).to_pytimedelta())
    assert freq.total > 0 and 0 <= freq.hits <= freq.total
    inter = serie_a[(serie_a["home_team"] == "Inter") | (serie_a["away_team"] == "Inter")].sort_values("date").tail(5)
    napoli = serie_a[(serie_a["home_team"] == "Napoli") | (serie_a["away_team"] == "Napoli")].sort_values("date").tail(5)
    sample = pd.concat([inter, napoli]).drop_duplicates(subset=["date", "home_team", "away_team"])
    assert freq.total == len(sample)
    assert freq.hits == int(((sample["home_goals"] + sample["away_goals"]) > 2.5).sum())
    assert freq.sentence().startswith(f"Nelle ultime {freq.total} partite delle due squadre «Over 2.5 gol» è uscito")
    assert recent_form.recent_frequency("1", "Inter", "Napoli", serie_a, as_of=last_day) is None   # dipende dal lato
    small = recent_form.Frequency("Goal", 2, 4, 2)
    assert small.small and "campione piccolo" in small.sentence()


def test_saved_slip_status_counts_and_free_results(serie_a):
    first = serie_a.sort_values("date").iloc[0]
    ordered = serie_a.sort_values("date")
    second = ordered[(ordered["home_goals"] + ordered["away_goals"]) > 0].iloc[1]
    legs = [Leg("a", f"{first['home_team']} - {first['away_team']}", ("Over 0.5 gol",), 0.9, 1.08),
            Leg("b", f"{second['home_team']} - {second['away_team']}", ("Under 0.5 gol",), 0.05, 12.0),
            Leg("c", "Squadra X - Squadra Y", ("Goal",), 0.5, 1.9)]
    slip = slips.make_slip("Multipla 31/10", legs, {"a": first["date"].date(), "b": second["date"].date(), "c": date(2026, 10, 31)})
    assert slip.score() == "0/3" and slip.state == slips.PENDING and slip.odds == pytest.approx(1.08 * 12 * 1.9, abs=0.01)
    changed = slips.update_from_results(slip, serie_a)
    assert changed == 2 and [leg["status"] for leg in slip.legs] == [slips.WON, slips.LOST, slips.PENDING]
    assert slip.score() == "1/3" and slip.state == slips.LOST and "football-data.co.uk" in slip.legs[0]["note"]
    restored = slips.Slip.from_dict(json.loads(json.dumps(slip.to_dict())))
    assert restored.to_dict() == slip.to_dict() and restored.id == slip.id
    rows = slips.sheet_rows([slip])
    assert len(rows) == 3 and rows[0]["Stato"] == "persa (1/3)"
    frame = slips.registry_like(slip)
    assert list(frame["Esito"].isna()) == [False, False, True]


def test_day_file_keeps_saved_slips():
    from match_memory import SLIPS_SHEET, day_workbook, read_day_workbook

    slip = slips.make_slip("Prova", [Leg("a", "A - B", ("Goal",), 0.5, 1.9)], {"a": date(2026, 10, 31)})
    content = day_workbook([], [], slips=[slip.to_dict()])
    from io import BytesIO
    from openpyxl import load_workbook

    sheet = list(load_workbook(BytesIO(content))[SLIPS_SHEET].iter_rows(values_only=True))
    assert sheet[1][0] == "Prova" and sheet[1][5] == "Goal"
    assert read_day_workbook(content).slips == [slip.to_dict()]


@pytest.fixture(scope="module")
def napoli():
    data = json.loads((SOURCES / "understat_players_Napoli_2026.json").read_text(encoding="utf-8"))
    return scorers.players_from_understat(data["players"])


def test_scorer_share_is_shrunk_and_probability_comes_from_the_simulated_goals(napoli):
    team_xg90 = scorers.team_xg_per_90(napoli)
    assert team_xg90 == pytest.approx(sum(player.xg for player in napoli) / max(player.games for player in napoli))
    striker = next(player for player in napoli if player.name == "Rasmus Højlund")
    raw = striker.xg / (team_xg90 * striker.minutes / 90)
    shrunk = scorers.share(striker, team_xg90)
    assert min(raw, scorers.PRIOR_SHARE) <= shrunk <= max(raw, scorers.PRIOR_SHARE)
    result = simulated(1850, 1750)
    distribution = scorers.goal_distribution(result, "home")
    assert sum(distribution.values()) == pytest.approx(1.0, abs=0.01)
    rows = scorers.scorer_rows(napoli, distribution)
    top = rows[0]
    assert rows == sorted(rows, key=lambda row: -row["Segna %"]) and top["Titolare"] == "non confermato"
    player = next(p for p in napoli if p.name == top["Giocatore"])
    # Fase P: con la sezione «marcatori» della configurazione vale la versione v2 (npxG, ruolo, rigorista).
    params = scorers.load_params()
    assert params is not None and params.verified
    team_npxg90 = sum(p.npxg for p in napoli) / max(p.games for p in napoli)
    taker = scorers.penalty_taker(napoli)
    per_goal = scorers.share_v2(player, team_npxg90, params, taker) * min(90, player.minutes_per_game) / 90
    expected = 1 - sum(p * (1 - per_goal) ** g for g, p in distribution.items())
    assert top["Segna %"] == pytest.approx(expected * 100, abs=0.05)
    assert top["Rigorista"] == ("sì" if top["Giocatore"] == taker else "")
    assert all(row["Segna %"] < 60 for row in rows)
    confirmed = scorers.scorer_rows(napoli, distribution, confirmed=[top["Giocatore"]])
    assert confirmed[0]["Minuti attesi"] == 90 and confirmed[0]["Titolare"] == "confermato"
    assert "verificata" in scorers.calibration_note()


def test_scorer_market_margin_from_entered_odds():
    rows = [{"Giocatore": "A", "Segna %": 40.0}, {"Giocatore": "B", "Segna %": 20.0}]
    margin = scorers.market_margin(rows, {"A": 2.0, "B": 3.5})
    assert margin == pytest.approx((1 / 2.0 + 1 / 3.5) / 0.6 - 1)
    assert scorers.market_margin(rows, {}) is None
