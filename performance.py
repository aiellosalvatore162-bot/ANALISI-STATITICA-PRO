"""Pagina «Prestazioni»: misure calcolate dal file Excel del registro (solo colonne A–S).

Le colonne grigie T–Y del file (Profitto, CLV, Brier) non vengono lette: tutto
si ricalcola qui dalle colonne di input. Esiti e quote di chiusura mancanti si
possono ricavare dai CSV di football-data.co.uk (risultati, corner, gialli,
quote medie di chiusura 1X2 e Over/Under 2.5) senza mai sovrascrivere i valori
già presenti nel file.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
import math
import re
from typing import Any, Callable, Iterable

import numpy as np
import pandas as pd

from engine import SECOND_YELLOW_EXCLUDED_LEAGUES
from prediction_log import COLUMNS, MARKETS, PROBABILITY_COLUMNS, PUSH_SELECTIONS, selection_id

MIN_SELECTIONS = 100
BANDS = (0.0, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0001)
COMPONENT_LABELS = {"Prob_mercato": "Mercato", "Prob_modello": "Modello", "Prob_stile": "Stile",
                    "Prob_Claude": "Claude", "Prob_finale": "Finale"}
# Confronti a coppie: ogni componente contro il mercato, stile e Claude anche contro il solo modello.
COMPARISONS = (("Prob_modello", "Prob_mercato"), ("Prob_stile", "Prob_mercato"), ("Prob_stile", "Prob_modello"),
               ("Prob_Claude", "Prob_mercato"), ("Prob_Claude", "Prob_modello"), ("Prob_finale", "Prob_mercato"))
# Quote medie di chiusura di football-data.co.uk (colonne AvgC*), nello schema di tools.dataset.
CLOSING_COLUMNS = {"1": "close_home", "X": "close_draw", "2": "close_away",
                   "Over 2.5 gol": "close_over25", "Under 2.5 gol": "close_under25"}
NUMERIC = (*PROBABILITY_COLUMNS, "Quota_presa", "Puntata", "Quota_chiusura")


class RegistryError(ValueError):
    pass


@dataclass
class Registry:
    frame: pd.DataFrame
    sheet: str
    warnings: list[str] = field(default_factory=list)
    duplicates: int = 0


# --- Lettura del file ----------------------------------------------------------------------------

def _number(value: Any) -> float | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _outcome(value: Any) -> float | None:
    if isinstance(value, str):
        text = value.strip().lower()
        if text in ("vinta", "v", "si", "sì", "true"):
            return 1.0
        if text in ("persa", "p", "no", "false"):
            return 0.0
    number = _number(value)
    return number if number in (0.0, 1.0) else None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _time(value: Any) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value or "").strip()


def read_registry(source: bytes | str | BytesIO) -> Registry:
    """Legge il foglio «Registro» (o il primo foglio con l'intestazione ID…Note), colonne A–S."""
    from openpyxl import load_workbook

    stream = BytesIO(source) if isinstance(source, bytes) else source
    try:
        workbook = load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:  # file non xlsx o danneggiato
        raise RegistryError(f"File Excel non leggibile: {exc}") from exc
    sheets = (["Registro"] if "Registro" in workbook.sheetnames else []) + \
        [name for name in workbook.sheetnames if name != "Registro"]
    for name in sheets:
        rows = workbook[name].iter_rows(min_row=1, max_col=len(COLUMNS), values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
        if "ID" in header and "Esito" in header:
            break
    else:
        raise RegistryError("Nessun foglio con l'intestazione del registro (ID … Note) nelle colonne A–S.")
    missing = [column for column in COLUMNS if column not in header]
    if missing:
        raise RegistryError("Colonne mancanti nel foglio «" + name + "» (A–S): " + ", ".join(missing) + ".")
    index = {column: header.index(column) for column in COLUMNS}
    warnings: list[str] = []
    records = []
    for excel_row, values in enumerate(rows, start=2):
        values = list(values) + [None] * (len(COLUMNS) - len(values))
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        record = {column: values[index[column]] for column in COLUMNS}
        record["_riga"] = excel_row
        records.append(record)
    workbook.close()
    frame = pd.DataFrame(records, columns=[*COLUMNS, "_riga"])
    if frame.empty:
        return Registry(frame.assign(Esito=pd.Series(dtype=float)), name, ["Il registro non contiene righe."])
    for column in NUMERIC:
        frame[column] = frame[column].map(_number).astype(float)
    percent = frame[list(PROBABILITY_COLUMNS)] > 1.0
    if percent.any().any():
        warnings.append("Alcune probabilità erano in percentuale (es. 62 invece di 0,62): convertite in decimali.")
        frame[list(PROBABILITY_COLUMNS)] = frame[list(PROBABILITY_COLUMNS)].mask(percent, frame[list(PROBABILITY_COLUMNS)] / 100.0)
    bad_outcome = frame["Esito"].notna() & frame["Esito"].map(_outcome).isna() & (frame["Esito"].astype(str).str.strip() != "")
    if bad_outcome.any():
        warnings.append(f"{int(bad_outcome.sum())} esiti non validi (ammessi solo 1 o 0): ignorati.")
    frame["Esito"] = frame["Esito"].map(_outcome).astype(float)
    frame["Data"] = frame["Data"].map(_date)
    frame["Ora"] = frame["Ora"].map(_time)
    for column in ("Competizione", "Partita", "Mercato", "Selezione", "Tipo", "Versione_modello", "Note", "ID"):
        frame[column] = frame[column].map(lambda value: "" if value is None or pd.isna(value) else str(value).strip())
    unknown = sorted(set(frame["Mercato"]) - set(MARKETS) - {""})
    if unknown:
        warnings.append("Mercati con nome non previsto dal Riepilogo: " + ", ".join(unknown) + ".")
    no_id = frame["ID"] == ""
    frame.loc[no_id, "ID"] = [selection_id(row.Data or "", row.Partita, row.Mercato, row.Selezione)
                              for row in frame[no_id].itertuples()]
    duplicates = int(frame["ID"].duplicated(keep="last").sum())
    if duplicates:
        warnings.append(f"{duplicates} righe duplicate (stesso ID): conteggiata solo l'ultima.")
    frame = frame.drop_duplicates("ID", keep="last").reset_index(drop=True)
    return Registry(frame, name, warnings, duplicates)


# --- Misure ---------------------------------------------------------------------------------------

def _settled(frame: pd.DataFrame, column: str = "Prob_finale") -> pd.DataFrame:
    return frame[frame["Esito"].notna() & frame[column].notna()]


def brier(probabilities, outcomes) -> float:
    return float(np.mean((np.asarray(probabilities, float) - np.asarray(outcomes, float)) ** 2))


def log_loss(probabilities, outcomes) -> float:
    p = np.clip(np.asarray(probabilities, float), 1e-6, 1 - 1e-6)
    y = np.asarray(outcomes, float)
    return float(-np.mean(y * np.log(p) + (1 - y) * np.log(1 - p)))


def _market_order(markets: Iterable[str]) -> list[str]:
    known = [market for market in MARKETS if market in set(markets)]
    return known + sorted(set(markets) - set(known))


def calibration_table(frame: pd.DataFrame, column: str = "Prob_finale", bands=BANDS) -> pd.DataFrame:
    """Per mercato e fascia di probabilità: selezioni, probabilità media prevista e frequenza reale."""
    data = _settled(frame, column)
    rows = []
    for market in _market_order(data["Mercato"]):
        subset = data[data["Mercato"] == market]
        for low, high in zip(bands[:-1], bands[1:]):
            band = subset[(subset[column] >= low) & (subset[column] < high)]
            if band.empty:
                continue
            predicted, observed = float(band[column].mean()), float(band["Esito"].mean())
            rows.append({"Mercato": market, "Fascia": f"{low:.0%}–{min(high, 1.0):.0%}", "Selezioni": len(band),
                         "Prob. prevista": predicted, "Frequenza reale": observed, "Scarto": observed - predicted})
    return pd.DataFrame(rows, columns=["Mercato", "Fascia", "Selezioni", "Prob. prevista", "Frequenza reale", "Scarto"])


def scores_over_time(frame: pd.DataFrame, freq: str = "M", column: str = "Prob_finale") -> pd.DataFrame:
    """Brier e log-loss per periodo (``M`` mese, ``W`` settimana) e cumulati."""
    data = _settled(frame, column)
    data = data[data["Data"].notna()].sort_values(["Data", "Ora"])
    columns = ["Periodo", "Selezioni", "Brier", "Log-loss", "Brier cumulato", "Log-loss cumulato"]
    if data.empty:
        return pd.DataFrame(columns=columns)
    periods = pd.to_datetime(data["Data"]).dt.to_period(freq).dt.start_time
    rows, seen = [], []
    for start, group in data.groupby(periods, sort=True):
        seen.append(group)
        history = pd.concat(seen)
        rows.append({"Periodo": start.date(), "Selezioni": len(group),
                     "Brier": brier(group[column], group["Esito"]), "Log-loss": log_loss(group[column], group["Esito"]),
                     "Brier cumulato": brier(history[column], history["Esito"]),
                     "Log-loss cumulato": log_loss(history[column], history["Esito"])})
    return pd.DataFrame(rows, columns=columns)


def with_clv(frame: pd.DataFrame) -> pd.DataFrame:
    """CLV = quota presa / quota di chiusura − 1 (come la colonna U del file)."""
    valid = frame["Quota_presa"].gt(1.0) & frame["Quota_chiusura"].gt(1.0)
    return frame.assign(CLV=(frame["Quota_presa"] / frame["Quota_chiusura"] - 1.0).where(valid))


def with_profit(frame: pd.DataFrame) -> pd.DataFrame:
    """Profitto = puntata · (quota · esito − 1), solo con esito, puntata e quota (colonna T del file)."""
    valid = frame["Esito"].notna() & frame["Puntata"].notna() & frame["Quota_presa"].notna()
    return frame.assign(Profitto=(frame["Puntata"] * (frame["Quota_presa"] * frame["Esito"] - 1.0)).where(valid))


def clv_by_market(frame: pd.DataFrame) -> pd.DataFrame:
    data = with_clv(frame)
    data = data[data["CLV"].notna()]
    rows = [{"Mercato": market, "Selezioni con chiusura": int((data["Mercato"] == market).sum()),
             "CLV medio": float(data.loc[data["Mercato"] == market, "CLV"].mean()),
             "Quote battute": float((data.loc[data["Mercato"] == market, "CLV"] > 0).mean())}
            for market in _market_order(data["Mercato"])]
    return pd.DataFrame(rows, columns=["Mercato", "Selezioni con chiusura", "CLV medio", "Quote battute"])


def bankroll(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Riepilogo per tipo (Reale / A carta) e saldo cumulato nel tempo, separati."""
    data = with_clv(with_profit(frame))
    played = data[data["Profitto"].notna()].sort_values(["Data", "Ora"], na_position="last")
    summary = []
    for kind in ["Reale", "A carta"] + sorted(set(played["Tipo"]) - {"Reale", "A carta"}):
        subset = played[played["Tipo"] == kind]
        if subset.empty and kind not in ("Reale", "A carta"):
            continue
        stake = float(subset["Puntata"].sum())
        clv = data.loc[(data["Tipo"] == kind) & data["CLV"].notna(), "CLV"]
        summary.append({"Tipo": kind or "(vuoto)", "Selezioni con esito": len(subset), "Puntata totale": stake,
                        "Profitto": float(subset["Profitto"].sum()), "ROI": float(subset["Profitto"].sum() / stake) if stake else None,
                        "CLV medio": float(clv.mean()) if len(clv) else None})
    timeline = played.assign(Saldo=played.groupby("Tipo")["Profitto"].cumsum())[["Data", "Ora", "Tipo", "Partita", "Selezione", "Profitto", "Saldo"]]
    return pd.DataFrame(summary), timeline.reset_index(drop=True)


def component_comparison(frame: pd.DataFrame) -> pd.DataFrame:
    """Brier e log-loss di ogni componente e confronto a coppie sulle stesse selezioni.

    «Migliora» solo se la differenza media di Brier è negativa oltre due errori
    standard: con poche selezioni il giudizio resta «non distinguibile».
    """
    settled = frame[frame["Esito"].notna()]
    rows = []
    for column, label in COMPONENT_LABELS.items():
        data = settled[settled[column].notna()]
        rows.append({"Confronto": label, "Selezioni": len(data),
                     "Brier": brier(data[column], data["Esito"]) if len(data) else None,
                     "Log-loss": log_loss(data[column], data["Esito"]) if len(data) else None,
                     "Δ Brier": None, "Giudizio": ""})
    for column, reference in COMPARISONS:
        data = settled[settled[column].notna() & settled[reference].notna()]
        if data.empty:
            continue
        difference = (data[column] - data["Esito"]) ** 2 - (data[reference] - data["Esito"]) ** 2
        mean = float(difference.mean())
        error = float(difference.std(ddof=1) / math.sqrt(len(data))) if len(data) > 1 else math.inf
        verdict = "migliora davvero" if mean + 2 * error < 0 else "peggiora" if mean - 2 * error > 0 else "non distinguibile"
        rows.append({"Confronto": f"{COMPONENT_LABELS[column]} vs {COMPONENT_LABELS[reference]}", "Selezioni": len(data),
                     "Brier": brier(data[column], data["Esito"]), "Log-loss": log_loss(data[column], data["Esito"]),
                     "Δ Brier": mean, "Giudizio": verdict})
    return pd.DataFrame(rows, columns=["Confronto", "Selezioni", "Brier", "Log-loss", "Δ Brier", "Giudizio"])


def edge_verdicts(frame: pd.DataFrame, min_selections: int = MIN_SELECTIONS) -> pd.DataFrame:
    """Giudizio per mercato con la stessa regola del foglio Riepilogo (≥100 selezioni con esito e CLV medio > 0)."""
    data = with_clv(with_profit(frame))
    rows = []
    for market in _market_order(set(data["Mercato"]) | set(MARKETS)):
        subset = data[data["Mercato"] == market]
        settled = subset[subset["Esito"].notna()]
        clv = subset["CLV"].dropna()
        stake = float(settled.loc[settled["Profitto"].notna(), "Puntata"].sum())
        count = len(settled)
        if count < min_selections:
            verdict = f"Non ancora dimostrato (servono {min_selections}, ne hai {count})"
        elif len(clv) and clv.mean() > 0:
            verdict = "VANTAGGIO DIMOSTRATO"
        else:
            verdict = "Nessun vantaggio (CLV non positivo)"
        rows.append({"Mercato": market, "Selezioni con esito": count,
                     "% vinte": float(settled["Esito"].mean()) if count else None,
                     "Prob. media finale": float(settled["Prob_finale"].mean()) if settled["Prob_finale"].notna().any() else None,
                     "Brier finale": brier(settled["Prob_finale"].dropna(), settled.loc[settled["Prob_finale"].notna(), "Esito"]) if settled["Prob_finale"].notna().any() else None,
                     "CLV medio": float(clv.mean()) if len(clv) else None,
                     "ROI": float(settled["Profitto"].sum() / stake) if stake else None,
                     "Giudizio": verdict})
    return pd.DataFrame(rows)


# --- Esiti e quote di chiusura dalle fonti --------------------------------------------------------

def split_match(match: str) -> tuple[str, str] | None:
    parts = re.split(r"\s+(?:-|–|vs\.?|v)\s+", match.strip(), maxsplit=1, flags=re.IGNORECASE)
    return (parts[0].strip(), parts[1].strip()) if len(parts) == 2 and all(parts) else None


def _needs(selection: str) -> list[str]:
    needed = ["home_goals", "away_goals"]
    if " 1T" in selection or " 2T" in selection or "BTTS" in selection:
        needed += ["home_ht_goals", "away_ht_goals"]
    if "corner" in selection:
        needed += ["home_corners", "away_corners"]
    if "cartellini" in selection:
        needed += ["home_yellows", "away_yellows", "home_reds", "away_reds"]
    return needed


def settle(selection: str, result: pd.Series | dict, card_rule=None) -> int | None:
    """Esito (1/0) di una selezione dal risultato reale, con le stesse definizioni degli eventi del motore.

    I cartellini sono punti con la regola della configurazione (``engine.CardRule``,
    default Eurobet: giallo 1, rosso 2). ``None`` se mancano i dati necessari o
    la selezione non è riconosciuta.
    """
    from engine import DEFAULT_CARD_RULE, _event_catalog, _team_multigol_ranges, card_points

    legs = selection.split(" + ")
    needed = {column for leg in legs for column in _needs(leg)}
    values = {}
    for column in needed | {"home_goals", "away_goals"}:
        value = result.get(column)
        if value is None or pd.isna(value):
            return None
        values[column] = np.array([int(value)])
    zero = np.array([0])
    if "home_reds" in values:
        rule = card_rule or DEFAULT_CARD_RULE
        if rule.double_yellow_extra:
            return None  # i dati non distinguono i rossi da doppia ammonizione
        values["home_yellows"] = card_points(values["home_yellows"], values["home_reds"], rule)
        values["away_yellows"] = card_points(values["away_yellows"], values["away_reds"], rule)
    first_home, first_away = values.get("home_ht_goals", zero), values.get("away_ht_goals", zero)
    events = _event_catalog(
        values["home_goals"], values["away_goals"], first_home, first_away,
        values["home_goals"] - first_home, values["away_goals"] - first_away,
        values.get("home_corners", zero), values.get("away_corners", zero),
        values.get("home_yellows", zero), values.get("away_yellows", zero),
        _team_multigol_ranges({leg: 2.0 for leg in legs}),
    )
    if any(leg not in events for leg in legs):
        return None
    return int(all(bool(events[leg][0]) for leg in legs))


def find_result(results: pd.DataFrame, match: str, day: date | None, *, before: int = 1, after: int = 10) -> pd.Series | None:
    """La partita giocata nella finestra [Data − before, Data + after] con entrambe le squadre riconosciute."""
    from datasources.names import match_team

    teams = split_match(match)
    if teams is None or day is None or results.empty:
        return None
    window = results[(results["date"] >= pd.Timestamp(day - timedelta(days=before)))
                     & (results["date"] <= pd.Timestamp(day + timedelta(days=after)))]
    if window.empty:
        return None
    home = match_team(teams[0], window["home_team"].unique())
    away = match_team(teams[1], window["away_team"].unique())
    if home is None or away is None:
        return None
    found = window[(window["home_team"] == home) & (window["away_team"] == away)]
    return None if found.empty else found.sort_values("date").iloc[0]


@dataclass
class FillReport:
    frame: pd.DataFrame
    outcomes: int = 0
    closings: int = 0
    pushes: int = 0
    not_found: list[str] = field(default_factory=list)
    changes: list[dict] = field(default_factory=list)


def fill_from_results(frame: pd.DataFrame, results: pd.DataFrame, *, source: str = "football-data.co.uk",
                      today: date | None = None) -> FillReport:
    """Compila solo Esito e Quota_chiusura vuoti; ogni valore aggiunto è annotato in Note con fonte e data."""
    today = today or date.today()
    filled = frame.copy()
    report = FillReport(filled)
    for position, row in filled.iterrows():
        need_outcome = pd.isna(row["Esito"]) and "rimborsata" not in row["Note"]
        closing_column = CLOSING_COLUMNS.get(row["Selezione"])
        need_closing = pd.isna(row["Quota_chiusura"]) and closing_column is not None
        if not (need_outcome or need_closing):
            continue
        match = find_result(results, row["Partita"], row["Data"])
        if match is None:
            if need_outcome:
                report.not_found.append(row["Partita"])
            continue
        added = {}
        details = []
        if need_outcome and row["Selezione"] in PUSH_SELECTIONS and match["home_goals"] == match["away_goals"]:
            # DNB con pareggio: puntata restituita, Esito vuoto, esclusa da Brier e ROI.
            details.append("rimborsata (pareggio): Esito lasciato vuoto")
            report.pushes += 1
        elif need_outcome:
            outcome = settle(row["Selezione"], match)
            if outcome is not None:
                added["Esito"] = float(outcome)
                report.outcomes += 1
                if ("cartellini" in row["Selezione"] and match.get("league") in SECOND_YELLOW_EXCLUDED_LEAGUES
                        and match["home_reds"] + match["away_reds"] > 0):
                    details.append("cartellini: dati inglesi/scozzesi senza il primo giallo della doppia ammonizione, "
                                   "un'eventuale doppia ammonizione conta 2 invece di 3")
        if need_closing and pd.notna(match.get(closing_column)) and float(match[closing_column]) > 1.0:
            added["Quota_chiusura"] = float(match[closing_column])
            report.closings += 1
        if not (added or details):
            continue
        what = " e ".join({"Esito": "esito", "Quota_chiusura": "quota di chiusura media"}[key] for key in added) or "risultato"
        note = "; ".join(details + [f"{what} da {source}"]) + f" ({match['date']:%d/%m/%Y} {int(match['home_goals'])}-{int(match['away_goals'])}), compilato il {today:%d/%m/%Y}"
        added["Note"] = f"{row['Note']}; {note}" if row["Note"] else note
        for column, value in added.items():
            filled.at[position, column] = value
        report.changes.append({"_riga": row["_riga"], "ID": row["ID"], **added})
    report.not_found = sorted(set(report.not_found))
    return report


def seasons_and_leagues(frame: pd.DataFrame, competitions: dict[str, dict]) -> tuple[list[int], list[str], list[str]]:
    """Stagioni (anno d'inizio) e codici football-data.co.uk da scaricare per le righe senza esito o chiusura."""
    from datasources.sources import season_start

    open_rows = frame[frame["Esito"].isna() | frame["Quota_chiusura"].isna()]
    days = [day for day in open_rows["Data"] if day is not None]
    seasons = sorted({season_start(day) for day in days})
    names = {name.lower(): codes["footballdata"] for name, codes in competitions.items()}
    names.update({codes["footballdata"].lower(): codes["footballdata"] for codes in competitions.values()})
    requested = {str(value).strip() for value in open_rows["Competizione"] if str(value).strip()}
    leagues = sorted({names[value.lower()] for value in requested if value.lower() in names})
    unknown = sorted(value for value in requested if value.lower() not in names)
    if not leagues or any(not str(value).strip() for value in open_rows["Competizione"]):
        leagues = sorted({codes["footballdata"] for codes in competitions.values()})
    return seasons, leagues, unknown


def collect_results(frame: pd.DataFrame, fetch_season: Callable[[str, int], pd.DataFrame],
                    competitions: dict[str, dict]) -> tuple[pd.DataFrame, list[str]]:
    seasons, leagues, unknown = seasons_and_leagues(frame, competitions)
    warnings = [f"Competizioni non coperte da football-data.co.uk (esiti da inserire a mano): {', '.join(unknown)}."] if unknown else []
    frames = [fetch_season(league, season) for league in leagues for season in seasons]
    frames = [data for data in frames if data is not None and not data.empty]
    return (pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()), warnings


def updated_workbook(original: bytes, sheet: str, changes: list[dict]) -> bytes:
    """Il file dell'utente con i soli valori aggiunti (Q, R, S); formule e altri fogli restano intatti."""
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(original))
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1][:len(COLUMNS)]]
    position = {column: header.index(column) + 1 for column in COLUMNS}
    ids = {}
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=position["ID"]).value
        if value:
            ids.setdefault(str(value).strip(), []).append(row)
    for change in changes:
        # Le righe duplicate (stesso ID) ricevono lo stesso valore.
        for row in sorted(set(ids.get(change["ID"], [])) | {change["_riga"]}):
            for column in ("Quota_chiusura", "Esito"):
                cell = worksheet.cell(row=row, column=position[column])
                if column in change and cell.value in (None, ""):
                    cell.value = int(change[column]) if column == "Esito" else change[column]
        if "Note" in change:
            worksheet.cell(row=change["_riga"], column=position["Note"]).value = change["Note"]
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
