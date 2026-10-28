"""Blocco FORMA del report dell'Analista: lettura tollerante, nessun effetto sul motore, colonna «Forma»."""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

import analyst
import dashboard
from analyst import Target, match_blocks, parse_report
from datasources.cache import SourceCache
from engine import config_from_input, simulate_match
from match_form import parse_form
from match_memory import day_workbook, read_day_workbook
from parser import parse_match_input
from prediction_log import COLUMNS, FORM_COLUMN, FORM_COLUMN_INDEX, PredictionLog, to_excel
from test_engine import VALID

BASE = """PARTITA: Genoa - Venezia | Serie A | 26/10/2026 15:00
Il Genoa recupera il centravanti.
{form}CORREZIONI_CLAUDE:
gol_casa: 1.05 | motivo: rientro del bomber | fonte: https://www.pianetagenoa1893.net/rientro
cartellini: 1.08 | motivo: arbitro severo | fonte: https://www.aia-figc.it/designazioni
confidenza: media
commento: Partita equilibrata.

CHIAVI_MATCH:
duello: centravanti Genoa contro centrali del Venezia | mercato: gol | effetto: più occasioni | fonte: https://esempio.it/duello
"""

FORM = """FORMA:
casa: V-V-P-N-V | gol_fatti_subiti: 9-4 | in_casa: V-N-V | fonte: https://www.legaseriea.it/genoa
ospite: P-N-V-P-P | gol_fatti_subiti: 5-8 | in_trasferta: P-P-N | fonte: https://www.legaseriea.it/venezia
nota: Genoa imbattuto in casa da tre partite.

"""

FORM_AFTER = BASE.replace("{form}", "") + "\n" + FORM  # il blocco può stare anche dopo le chiavi


def _block(text):
    return parse_report(text)[0]


def test_form_block_is_read_when_present():
    block = _block(BASE.replace("{form}", FORM))
    form = block.form
    assert form.home.results == ["V", "V", "P", "N", "V"] and form.home.goals == (9, 4) and form.home.venue == ["V", "N", "V"]
    assert form.away.results == ["P", "N", "V", "P", "P"] and form.away.source == "https://www.legaseriea.it/venezia"
    assert form.note == "Genoa imbattuto in casa da tre partite." and not form.warnings
    assert form.summary() == ("Casa V-V-P-N-V, gol 9-4, in casa V-N-V | Ospite P-N-V-P-P, gol 5-8, fuori P-P-N | "
                              "Genoa imbattuto in casa da tre partite.")
    assert _block(FORM_AFTER).form.summary() == form.summary()
    row = analyst.report_rows([block], {})[0]
    assert row["Forma"] == form.summary()


def test_without_form_block_everything_works_as_before():
    block = _block(BASE.replace("{form}", ""))
    assert block.form is None and not block.error
    assert analyst.report_rows([block], {})[0]["Forma"] == ""
    assert parse_form("nessun blocco qui") is None
    # Il report completo della fase Y non ha FORMA: invariato.
    assert all(item.form is None for item in analyst.parse_full(open("tests/fixtures/analista/report_completo.md").read()).matches)


def test_incomplete_form_keeps_valid_parts_and_never_fails():
    text = """FORMA:
casa: V-X-P | gol_fatti_subiti: nove | in_casa: V-N-V | fonte: https://esempio.it/a
ospite: P-N-V-P-P | gol_fatti_subiti: 5-8
nota: solo indicativa
"""
    form = parse_form(text)
    # Casa: sequenza e gol non validi scartati, resta l'andamento in casa; ospite senza fonte scartato.
    assert form.home.results == [] and form.home.goals is None and form.home.venue == ["V", "N", "V"]
    assert form.away is None and form.note == "solo indicativa"
    assert len(form.warnings) == 3
    assert parse_form("FORMA:\n\n") is None
    assert parse_form("FORMA:\ncasa: V-V | fonte: niente link\n") is None
    block = _block(BASE.replace("{form}", text + "\n"))
    assert not block.error and block.response is not None and block.form.note == "solo indicativa"


def test_form_never_changes_corrections_or_probabilities():
    plain, with_form = _block(BASE.replace("{form}", "")), _block(BASE.replace("{form}", FORM))
    late = _block(FORM_AFTER)
    for block in (with_form, late):
        assert block.corrections_text == plain.corrections_text and block.keys_text == plain.keys_text
        assert block.comment == plain.comment and block.response.warnings == plain.response.warnings
        assert block.response.adjustment() == plain.response.adjustment()
    config = config_from_input(parse_match_input(VALID), simulations=3000, seed=7)
    from dataclasses import replace

    first = simulate_match(replace(config, claude=plain.response.adjustment()))
    second = simulate_match(replace(config, claude=with_form.response.adjustment()))
    assert first["components"] == second["components"] and first["model"] == second["model"]
    assert first["financial"] == second["financial"]


def test_swapped_teams_swap_the_form_once():
    text = BASE.replace("Genoa - Venezia", "Venezia - Genoa").replace("{form}", FORM)
    blocks = parse_report(text)
    targets = [Target(0, "Genoa", "Venezia", date(2026, 10, 26))]
    for _ in range(2):  # abbinare di nuovo non deve scambiare di nuovo
        match_blocks(blocks, targets)
        assert blocks[0].swapped and blocks[0].form.home.results == ["P", "N", "V", "P", "P"]


def test_form_column_in_registry_and_excel(tmp_path):
    log = PredictionLog(SourceCache(tmp_path / "cache.sqlite", sleep=lambda _: None, min_interval={}))
    row = {column: None for column in COLUMNS} | {"ID": "a", "Data": "2026-10-26", "Partita": "Genoa - Venezia",
                                                     "Mercato": "1X2", "Selezione": "1", "Note": ""}
    log.save([{**row, FORM_COLUMN: "Casa V-V-P-N-V"}], "s")
    assert log.rows("s")[0][FORM_COLUMN] == "Casa V-V-P-N-V"
    log.save([{**row, "Esito": 1.0}], "s")  # riga riletta dalle sole A–S: la forma resta
    saved = log.rows("s")[0]
    assert saved[FORM_COLUMN] == "Casa V-V-P-N-V" and saved["Esito"] == 1.0
    sheet = load_workbook(BytesIO(to_excel([saved])))["Registro"]
    assert tuple(cell.value for cell in sheet[1])[:19] == COLUMNS
    assert sheet.cell(row=1, column=FORM_COLUMN_INDEX).value == "Forma"
    assert sheet.cell(row=2, column=FORM_COLUMN_INDEX).value == "Casa V-V-P-N-V"
    restored = read_day_workbook(day_workbook([saved], []))
    assert restored.registry_rows[0][FORM_COLUMN] == "Casa V-V-P-N-V"
    assert load_workbook(BytesIO(to_excel([row])))["Registro"].max_column == 19  # senza forma: solo A–S


def test_form_card_is_informative_and_escaped():
    form = parse_form(FORM.replace("Genoa imbattuto", "<b>Genoa</b> imbattuto"))
    html = dashboard.form_card("Genoa", "Venezia", form)
    assert "V-V-P-N-V" in html and "&lt;b&gt;Genoa&lt;/b&gt;" in html and "<b>Genoa" not in html
    assert "non cambia probabilità" in html
