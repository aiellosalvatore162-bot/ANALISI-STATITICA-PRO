"""Fase R: «Modifica o completa i dati» in Oggi: blocco parziale, anteprima, precedenza, ricalcolo di una partita, Excel."""

from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest

import manual_data
import today
from datasources.builder import render_block
from datasources.stats import MANUAL_SOURCE
from match_memory import day_workbook, read_day_workbook
from parser import PartialBlock, parse_match_input, parse_partial_input
from test_today_speed import FullWeb, _match, _provider
from today import analyze, prepare, problem_rows

ROOT = Path(__file__).parent


def _auto(tmp_path, home="Napoli", away="Roma"):
    provider = _provider(tmp_path / "c.sqlite", FullWeb())
    return provider.build(home, away, match_date=_match(home, away).commence_time.date())


def test_partial_block_reads_only_the_present_fields_and_ignores_comments():
    block = parse_partial_input("# CASA: Roma | Elo: 1\n# Attacco: 9\nCASA: Napoli | Elo: 1650\nARBITRO: Gialli medi 5,1\nFoo: 3")
    assert (block.home_name, block.away_name) == ("Napoli", None)
    assert block.home == {"elo": 1650.0} and block.away == {} and block.referee == {"yellow": 5.1}
    assert block.unrecognized == ("Foo: 3",) and not block.errors and not block.empty
    wrong = parse_partial_input("CASA: Napoli | Elo: 5000 | Attacco: 1,2 | Forma: w-d-x")
    assert wrong.home == {"attack": 1.2} and len(wrong.errors) == 2  # stessi controlli e messaggi del parser completo
    assert any("Elo: usare un rating tra 800 e 2400" in error for error in wrong.errors)
    assert parse_partial_input("# solo commenti").empty
    with pytest.raises(Exception):  # il parser completo resta esigente
        parse_match_input("CASA: Napoli | Elo: 1650")


def test_partial_block_changes_only_the_elo_and_keeps_source_of_the_rest(tmp_path):
    auto = _auto(tmp_path)
    block = parse_partial_input("CASA: Napoli | Elo: 1650")
    entry = manual_data.merge(None, manual_data.changes(block, auto, None), "Napoli - Roma")
    sheet = manual_data.apply_manual(auto, entry)
    assert sheet.home_fields["elo"].value == 1650.0 and sheet.home_fields["elo"].source == MANUAL_SOURCE
    assert sheet.home_fields["elo"].origin == "manuale"
    for key in ("attack", "defense", "form", "xg_for", "corners", "cards_for"):  # il resto è quello di prima, con la sua fonte
        assert sheet.home_fields[key] == auto.home_fields[key]
    assert sheet.away_fields == auto.away_fields and sheet.referee_fields == auto.referee_fields
    rows = {(row["Soggetto"], row["Campo"]): row for row in manual_data.field_table(sheet)}
    assert rows[("Napoli", "Elo")]["Origine"] == "manuale" and rows[("Napoli", "Elo")]["Fonte"] == MANUAL_SOURCE
    assert rows[("Napoli", "Attacco")]["Origine"] == "automatico" and "football-data" in rows[("Napoli", "Attacco")]["Fonte"]
    assert rows[("Roma", "Elo")]["Origine"] == "proxy/stima" and "Elo da risultati" in rows[("Roma", "Elo")]["Fonte"]


def test_preview_says_replaces_completes_same_and_unrecognized(tmp_path):
    auto = _auto(tmp_path)
    old = auto.home_fields["elo"]
    text = f"CASA: Napoli | Elo: 1650 | Attacco: {auto.home_fields['attack'].value} | Foo: 3\nARBITRO: Gialli medi 6"
    block = parse_partial_input(text)
    rows = {row["Campo"]: row for row in manual_data.preview(block, auto)}
    elo = rows["Casa · Elo"]
    assert elo["Automatico"] == f"{old.value:.0f} ({old.source})" and elo["Blocco"] == "1650"
    assert elo["Cosa succede"] == manual_data.REPLACES
    assert rows["Casa · Attacco"]["Cosa succede"] == manual_data.SAME
    assert rows["Foo: 3"]["Cosa succede"] == manual_data.UNRECOGNIZED
    assert "Arbitro · Gialli medi" in rows
    # Il valore uguale non diventa manuale; dopo l'applicazione, una nuova anteprima cita il manuale precedente.
    changed = manual_data.changes(block, auto, None)
    assert "attack" not in changed["home"] and changed["home"]["elo"] == 1650.0
    entry = manual_data.merge(None, changed, "Napoli - Roma")
    again = manual_data.preview(parse_partial_input("CASA: Napoli | Elo: 1700"), auto, entry)[0]
    assert again["Cosa succede"] == f"{manual_data.REPLACES} (manuale precedente: 1650)"
    cup = _provider(tmp_path / "cup.sqlite", FullWeb()).build("Napoli", "FC Porto", match_date=_match("a", "b").commence_time.date())
    assert manual_data.preview(parse_partial_input("CASA: Napoli | Elo: 1650"), cup)[0]["Cosa succede"] == manual_data.COMPLETES
    assert manual_data.preview(parse_partial_input("CASA: Napoli | Elo: 1650"), cup)[0]["Automatico"] == "MANCANTE"


def test_teams_are_checked_with_match_team():
    block = parse_partial_input("CASA: SSC Napoli | Elo: 1650")
    assert manual_data.check_teams(block, "Napoli", "Roma") is None
    wrong = manual_data.check_teams(parse_partial_input("CASA: Juventus | Elo: 1700"), "Napoli", "Roma")
    assert "non corrisponde" in wrong and "nulla applicato" in wrong
    swapped = manual_data.check_teams(parse_partial_input("CASA: Roma | Elo: 1700"), "Napoli", "Roma")
    assert "squadra ospite" in swapped
    assert "non indica la squadra" in manual_data.check_teams(PartialBlock(None, None, {"elo": 1600.0}), "Napoli", "Roma")
    assert manual_data.check_teams(parse_partial_input("ARBITRO: Gialli medi 5"), "Napoli", "Roma") is None


def test_manual_values_win_over_a_new_prepare_and_analyze_until_restored(tmp_path):
    db = tmp_path / "c.sqlite"
    first = prepare(_match("Napoli", "Roma"), _provider(db, FullWeb()).build, {}, simulations=2000)
    entry = manual_data.merge(None, {"home": {"elo": 1650.0}}, "Napoli - Roma")
    again = analyze(_match("Napoli", "Roma"), _provider(db, FullWeb()).build, {}, simulations=2000, manual=entry)  # nuova sessione
    row = next(r for r in again.sources if r["Soggetto"] == "Napoli" and r["Campo"] == "Elo")
    assert (row["Valore"], row["Origine"], row["Fonte"]) == ("1650", "manuale", MANUAL_SOURCE)
    assert "Elo: 1650" in again.block and again.ok
    assert again.config.home.elo == 1650.0 and first.config.home.elo != 1650.0
    restored = analyze(_match("Napoli", "Roma"), _provider(db, FullWeb()).build, {}, simulations=2000, manual=None)
    row = next(r for r in restored.sources if r["Soggetto"] == "Napoli" and r["Campo"] == "Elo")
    assert row["Origine"] == "proxy/stima" and restored.config.home.elo == first.config.home.elo


def test_cup_with_missing_elo_is_listed_then_analyzed_once_completed(tmp_path):
    provider = _provider(tmp_path / "c.sqlite", FullWeb())
    match = _match("Napoli", "FC Porto")
    waiting = prepare(match, provider.build, {}, simulations=2000)
    assert waiting.config is None and not waiting.ok and waiting.auto_sheet is not None
    row = problem_rows([waiting])[0]
    assert (row["Fase"], row["Fonte"]) == ("dati", "Elo da risultati") and "Modifica o completa i dati" in row["Cosa fare"]
    block = parse_partial_input("CASA: Napoli | Elo: 1650\nOSPITE: FC Porto | Elo: 1700")
    entry = manual_data.merge(None, manual_data.changes(block, waiting.auto_sheet, None), match.label)
    done = analyze(match, None, {}, simulations=2000, manual=entry, auto_sheet=waiting.auto_sheet)  # nessuna rete
    assert done.ok and done.problem is None and done.config.away.elo == 1700.0
    only_one = manual_data.merge(None, {"home": {"elo": 1650.0}}, match.label)
    partial = prepare(match, None, {}, simulations=2000, manual=only_one, auto_sheet=waiting.auto_sheet)
    assert partial.config is None and partial.problem[1] == "Elo da risultati"  # l'ospite resta MANCANTE


def test_excel_roundtrip_keeps_manual_values_origin_and_source(tmp_path):
    auto = _auto(tmp_path)
    block = parse_partial_input("CASA: Napoli | Elo: 1650 | Forma: W-W-D\nARBITRO: Gialli medi 5,5")
    entry = manual_data.merge(None, manual_data.changes(block, auto, None), "Napoli - Roma")
    items = [{"id": "2026-10-25|Napoli - Roma", **entry}]
    content = day_workbook([], [], manual=items)
    restored = read_day_workbook(content)
    assert restored.manual == items
    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(content))["Manuali"]
    rows = [dict(zip(manual_data.MANUAL_COLUMNS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert {(r["Campo"], r["Valore"], r["Origine"], r["Fonte"]) for r in rows} == {
        ("Elo", "1650", "manuale", MANUAL_SOURCE), ("Forma", "W-W-D", "manuale", MANUAL_SOURCE),
        ("Gialli medi", "5.5", "manuale", MANUAL_SOURCE)}
    assert manual_data.apply_manual(auto, restored.manual[0]) == manual_data.apply_manual(auto, entry)


# --- Flusso nell'interfaccia (AppTest): niente rete, un solo ricalcolo ------------------------------------

SCRIPT = '''
import sys
sys.path.insert(0, {root!r})
import streamlit as st
import test_manual as t
from today import analyze
from multiples import load_settings
from ui_manual import render_manual_section
if "ready" not in st.session_state:
    provider = t._provider({db!r}, t.FullWeb())
    st.session_state["today_analyses"] = [
        analyze(t._match(home, away), provider.build, {{}}, simulations=2000) for home, away in
        (("Napoli", "Roma"), ("Napoli", "FC Porto"))]
    st.session_state["ready"] = True
settings = load_settings()
render_manual_section(st.session_state["today_analyses"], settings)
'''


def test_ui_flow_complete_the_cup_recompute_one_match_and_restore(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest

    import ui_manual

    script = tmp_path / "pagina.py"
    script.write_text(SCRIPT.format(root=str(ROOT), db=str(tmp_path / "c.sqlite")))
    at = AppTest.from_file(str(script), default_timeout=120).run()
    assert not at.exception
    ok_match, cup = at.session_state["today_analyses"]
    assert ok_match.ok and not cup.ok
    labels = [e.label for e in at.expander]
    assert "MANCANTE" in labels[1] and "MANCANTE" not in labels[0]
    assert at.expander[1].proto.expanded and not at.expander[0].proto.expanded  # aperto solo quello con dati mancanti
    key = ui_manual.match_key(cup.match)
    text_key, apply_key = ui_manual._widget_key("manual_text", key), ui_manual._widget_key("manual_apply", key)

    at.text_area(key=text_key).input("CASA: Juventus | Elo: 1700").run()  # squadra sbagliata: avviso, niente applicazione
    assert any("non corrisponde" in warning.value for warning in at.warning) and at.button(key=apply_key).disabled

    calls = []
    original = today.simulate_match
    monkeypatch.setattr(today, "simulate_match", lambda config: calls.append(1) or original(config))
    at.text_area(key=text_key).input("# commento\nCASA: Napoli | Elo: 1650\nOSPITE: FC Porto | Elo: 1700").run()
    assert not at.button(key=apply_key).disabled
    assert any("completa MANCANTE" in str(frame.value) for frame in at.dataframe)  # anteprima
    at.button(key=apply_key).click().run()
    assert not at.exception and calls == [1]  # ricalcolata una sola partita, una sola simulazione
    after_ok, after_cup = at.session_state["today_analyses"]
    assert after_cup.ok and after_cup.config.home.elo == 1650.0 and after_ok is ok_match
    assert at.session_state["today_problems"] == [] and at.session_state["day_archive"][key]["manual"]
    assert "con dati a mano" in at.expander[1].label and not at.expander[1].proto.expanded
    # Scheda partita: la memoria della giornata conserva valore, origine e fonte di ogni campo.
    record = next(item for item in at.session_state["analyzed_matches"].values() if "Porto" in item["away"])
    fields = record["fields"]
    elo = next(row for row in fields if row["Soggetto"] == "Napoli" and row["Campo"] == "Elo")
    assert (elo["Valore"], elo["Origine"], elo["Fonte"]) == ("1650", "manuale", MANUAL_SOURCE)
    attack = next(row for row in fields if row["Soggetto"] == "Napoli" and row["Campo"] == "Attacco")
    assert attack["Origine"] == "automatico"

    restore_key = ui_manual._widget_key("manual_restore", key)
    at.button(key=restore_key).click().run()
    assert not at.exception and calls == [1]  # tornata «MANCANTE»: nessuna simulazione in più
    assert not at.session_state["today_analyses"][1].ok and not at.session_state["day_archive"][key]["manual"]
    assert at.expander[1].proto.expanded
