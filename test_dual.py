"""Doppio pronostico (fase U): scelte, fonti nel registro, classifiche, disaccordi e suggerimenti."""

from datetime import date, timedelta
from io import BytesIO

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

import dual
import engine
from match_memory import DUAL_SHEET, day_workbook, read_day_workbook
from multiples import Leg
from prediction_log import COLUMNS, SaveInfo, prediction_rows


def legs():
    return [Leg("0", "Napoli - Roma", ("1",), 0.62, 1.55), Leg("0", "Napoli - Roma", ("Over 1.5 gol",), 0.78, 1.18),
            Leg("0", "Napoli - Roma", ("1X",), 0.80, 1.10), Leg("0", "Napoli - Roma", ("Under 3.5 gol",), 0.70, 1.45),
            Leg("0", "Napoli - Roma", ("Over 2.5 gol", "Over 8.5 corner"), 0.95, 1.5)]


def test_simulator_pick_is_the_most_probable_single_that_is_not_a_trap():
    # 1X: EV −12% (trappola); Over 1.5 sotto la quota minima 1.20; il Combo Builder non è una singola.
    pick = dual.simulator_pick(legs(), trap_ev=-5.0, min_odds=1.20)
    assert pick.selections == ("Under 3.5 gol",)
    assert dual.simulator_pick(legs(), trap_ev=-50.0, min_odds=1.0).selections == ("1X",)
    assert dual.simulator_pick([], trap_ev=-5.0) is None


def test_unita_sources_and_strong_disagreement():
    sim, mine = dual.Pick("Under 3.5 gol", 0.70, 1.30), dual.Pick("2", 0.21, 3.90)
    choice = dual.DualChoice("2026-10-25|Napoli - Roma", "Napoli - Roma", sim, mine, mine)
    assert not choice.agree and choice.strong_disagreement(0.45)
    merged = dual.sources_by_selection(choice)
    assert merged == {"Under 3.5 gol": (["simulatore"], 1.30), "2": (["osservatore", "unita"], 3.90)}
    same = dual.DualChoice("x", "Napoli - Roma", sim, dual.Pick("Under 3.5 gol", 0.70, 1.35), dual.Pick("Under 3.5 gol", 0.70, 1.35))
    assert same.agree and not same.strong_disagreement(0.45)
    assert dual.sources_by_selection(same) == {"Under 3.5 gol": (["simulatore", "osservatore", "unita"], 1.35)}
    record = dual.choice_record(choice, date(2026, 10, 25), 0.45)
    assert record["strong"] and record["osservatore"] == {"selection": "2", "probability": 0.21, "odds": 3.90}


def test_observer_rows_keep_the_simulator_probability_and_the_source_in_note():
    team = lambda name, elo: engine.TeamParams(name, 1.1, 1.0, elo, "W-D-W-L-W", 1.5, 1.2, 5.5, 4.5, 2.0)
    result = engine.simulate_match(engine.MatchConfig(team("Napoli", 1850), team("Roma", 1800), 4.4, 24.0,
                                                      odds={"2": 3.9}, simulations=4_000))
    pick = dual.pick_for("2", result, 3.9)
    assert pick.probability == pytest.approx(result["components"]["finale"]["2"] / 100)
    info = SaveInfo(date(2026, 10, 25), "18:00", "Serie A", "A carta", "t", "segnalazione: turnover casa; " + dual.source_note(["osservatore"]))
    row = prediction_rows(result, ["2"], info, odds={"2": 3.9})[0]
    assert row["Prob_finale"] == pytest.approx(pick.probability, abs=1e-4) and row["Quota_presa"] == 3.9
    assert dual.sources_of(row["Note"]) == ["osservatore"]
    assert dual.sources_of("fonte: simulatore, unita; esito da football-data.co.uk") == ["simulatore", "unita"]
    assert dual.sources_of("esito da Claude") == []
    assert "2" in dual.observer_options(result, {"2": 3.9, "X": 0.0, "Over 4.5 gol": 9.0})
    assert "Over 4.5 gol" not in dual.observer_options(result, {"Over 4.5 gol": 9.0})  # non registrabile


def test_weekend_groups_friday_to_monday():
    saturday = date(2026, 10, 24)
    assert {dual.weekend_of(saturday + timedelta(days=offset)) for offset in (-1, 0, 1, 2)} == {saturday}
    assert dual.weekend_of(date(2026, 10, 27)) == date(2026, 10, 31)   # martedì: weekend successivo


def registry(weeks=1):
    """Ogni partita: simulatore e osservatore in disaccordo, più una in accordo."""
    rows = []
    for week in range(weeks):
        day = date(2026, 10, 3) + timedelta(days=7 * week)
        entries = [  # partita, selezione, mercato, prob simulatore, quota, esito, fonti
            ("A - B", "Over 2.5 gol", "Over 2.5", 0.60, 1.80, 1.0, ["simulatore"]),
            ("A - B", "Goal", "Goal/NoGoal", 0.40, 2.10, 0.0, ["osservatore", "unita"]),
            ("C - D", "1", "1X2", 0.55, 1.90, 0.0, ["simulatore"]),
            ("C - D", "X", "1X2", 0.25, 3.40, 1.0, ["osservatore", "unita"]),
            ("E - F", "Under 3.5 gol", "Under 3.5", 0.70, 1.35, 1.0, ["simulatore", "osservatore", "unita"]),
        ]
        for match, selection, market, probability, odds, outcome, sources in entries:
            rows.append({**{column: None for column in COLUMNS}, "ID": f"{day}|{match}|{market}|{selection}", "Data": day,
                         "Partita": match, "Mercato": market, "Selezione": selection, "Prob_finale": probability,
                         "Quota_presa": odds, "Esito": outcome, "Note": dual.source_note(sources)})
    frame = pd.DataFrame(rows)
    frame["Esito"] = frame["Esito"].astype(float)
    return frame


def test_leaderboards_by_source_market_and_weekend():
    frame = registry()
    table = dual.leaderboard(frame).set_index("Fonte")
    assert table.loc["Simulatore", "Selezioni"] == 3 and table.loc["Simulatore", "% vinte"] == pytest.approx(2 / 3)
    assert table.loc["Simulatore", "Attese"] == pytest.approx((0.60 + 0.55 + 0.70) / 3)
    assert table.loc["Simulatore", "ROI"] == pytest.approx(((1.80 - 1) + (0 - 1) + (1.35 - 1)) / 3)
    assert table.loc["Osservatore", "ROI"] == pytest.approx(((0 - 1) + (3.40 - 1) + (1.35 - 1)) / 3)
    assert table.loc["Osservatore", "Brier"] == pytest.approx(((0.40 - 0) ** 2 + (0.25 - 1) ** 2 + (0.70 - 1) ** 2) / 3)
    assert table.loc["Unita", "Selezioni"] == 3
    by_market = dual.leaderboard(frame, "Mercato")
    assert set(by_market["Mercato"]) == {"Over 2.5", "Goal/NoGoal", "1X2", "Under 3.5"}
    by_weekend = dual.leaderboard(registry(2), "Weekend")
    assert sorted(set(by_weekend["Weekend"])) == [date(2026, 10, 3), date(2026, 10, 10)]


def test_disagreements_tell_who_was_right():
    summary, details = dual.disagreements(registry())
    assert list(details["Ha ragione"]) == ["Simulatore", "Osservatore"] and len(details) == 2   # E - F in accordo
    assert dict(zip(summary["Ha ragione"], summary["Volte"])) == {"Simulatore": 1, "Osservatore": 1, "Entrambi": 0, "Nessuno": 0}


def test_suggestions_only_after_enough_weekends():
    lines, weekends = dual.suggestions(registry(3), min_weekends=4, min_selections=1)
    assert lines == [] and weekends == 3
    lines, weekends = dual.suggestions(registry(4), min_weekends=4, min_selections=1)
    assert weekends == 4
    assert any(line.startswith("1X2: su questo mercato rende meglio l'osservatore") for line in lines)
    assert not any(line.startswith("Under 3.5:") for line in lines)  # stessa selezione: nessun vincitore


def test_day_file_has_the_dual_sheet_and_restores_it():
    choice = dual.DualChoice("2026-10-25|Napoli - Roma", "Napoli - Roma", dual.Pick("Under 3.5 gol", 0.7, 1.3),
                             dual.Pick("2", 0.21, 3.9), dual.Pick("2", 0.21, 3.9))
    record = dual.choice_record(choice, date(2026, 10, 25), 0.45)
    content = day_workbook([], [], dual_choices=[record])
    sheet = list(load_workbook(BytesIO(content))[DUAL_SHEET].iter_rows(values_only=True))
    header = sheet[0]
    row = dict(zip(header, sheet[1]))
    assert row["Simulatore"] == "Under 3.5 gol" and row["Osservatore"] == "2" and row["Unita"] == "2"
    assert row["Quota (Osservatore)"] == 3.9 and row["Disaccordo forte"] == "sì"
    assert read_day_workbook(content).dual_choices == [record]
