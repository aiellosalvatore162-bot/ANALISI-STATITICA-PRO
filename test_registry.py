"""Test del registro Excel (fase H) e della pagina Prestazioni: solo file di esempio, nessuna rete."""

from dataclasses import replace
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import sqlite3

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import engine
from datasources.cache import SourceCache
from datasources.sources import COMPETITIONS
from parser import parse_match_input
from performance import (
    RegistryError, bankroll, brier, calibration_table, clv_by_market, collect_results, component_comparison,
    edge_verdicts, fill_from_results, log_loss, read_registry, scores_over_time, settle, split_match, updated_workbook,
)
from prediction_log import COLUMNS, MARKETS, PredictionLog, SaveInfo, market_of, model_version, prediction_rows, to_excel
from test_parser import VALID
from tools.dataset import clean_frames, frame_from_csv

FIXTURES = Path(__file__).resolve().parent / "tests" / "fixtures"
TEMPLATE = FIXTURES / "Registro_Previsioni.xlsx"
SAMPLE = FIXTURES / "registro_esempio.xlsx"
INFO = SaveInfo(date(2026, 10, 25), "20:45", "Serie A", "Reale", "loglineare · ancoraggio 0.9")


def _results() -> pd.DataFrame:
    body = (FIXTURES / "sources" / "footballdata_I1_2627.csv").read_bytes()
    return clean_frames([frame_from_csv(body, "I1", "2627")])


@pytest.fixture(scope="module")
def simulations():
    config = engine.config_from_input(parse_match_input(VALID), simulations=20_000, seed=7)
    config = replace(config, odds={"1": 1.9, "X": 3.5, "2": 4.2, "Over 2.5 gol": 1.8, "Under 2.5 gol": 2.0})
    base = engine.simulate_match(config)
    styled = engine.simulate_match(replace(config, tactical=engine.TacticalFactors(cards=1.1),
                                           tactical_switches=engine.TacticalSwitches(cards=True)))
    claude = engine.simulate_match(replace(config, claude=engine.ClaudeAdjustment(goals_home=0.95)))
    return config, base, styled, claude


# --- Registro di sessione ---------------------------------------------------------------------------

def test_columns_and_markets_match_the_user_template():
    workbook = load_workbook(TEMPLATE)
    header = [cell.value for cell in workbook["Registro"][1]]
    assert tuple(header[:19]) == COLUMNS
    summary = workbook["Riepilogo"]
    assert tuple(summary.cell(row=row, column=1).value for row in range(9, 9 + len(MARKETS))) == MARKETS
    assert len(MARKETS) == 29 and MARKETS[20] == "Multipla" and MARKETS[-1] == "Fasce cartellini"
    after = summary.cell(row=9 + len(MARKETS), column=1).value  # subito dopo l'elenco: vuota o la nota sui Brier
    assert after is None or str(after).startswith("Quale componente")


@pytest.mark.parametrize("selection, market", [
    ("1", "1X2"), ("X2", "Doppia chance"), ("12", "Doppia chance"), ("Casa DNB", "DNB"), ("Ospite DNB", "DNB"),
    ("Over 1.5 gol", "Over 1.5"), ("Under 1.5 gol", "Under 1.5"), ("Over 2.5 gol", "Over 2.5"), ("Under 2.5 gol", "Under 2.5"),
    ("Over 3.5 gol", "Over 3.5"), ("Under 3.5 gol", "Under 3.5"), ("Goal", "Goal/NoGoal"), ("No Goal", "Goal/NoGoal"),
    ("Multigol 2-4", "Multigol"), ("Multigol Casa 1-2", "Multigol casa"), ("Multigol Ospite 0-1", "Multigol ospite"),
    ("1 1T", "Primo tempo"), ("X 1T", "Primo tempo"), ("Over 0.5 gol 1T", "Over 0.5 1T"), ("Over 1.5 gol 1T", "Over 1.5 1T"),
    ("Over 9.5 corner", "Corner"), ("1 corner", "Corner"), ("Over 4.5 cartellini", "Cartellini"), ("1 + Over 1.5 gol", "Combo"),
    # Linee senza nome nella tabella del Riepilogo: non registrabili.
    ("Over 4.5 gol", "Over/Under 4.5+"), ("Somma gol 3", None), ("Under 0.5 gol", None), ("1X 1T", None), ("Goal 1T", None), ("Under 0.5 gol 1T", None),
    ("Casa DNB push", None), ("1-0/2-0/2-1", "Risultati esatti"), ("0-0", "Risultati esatti"), ("1-0/", None),
])
def test_market_names_are_fixed(selection, market):
    assert market_of(selection) == market


def test_rows_carry_every_component(simulations):
    config, base, styled, claude = simulations
    rows = prediction_rows(base, ["1", "Over 2.5 gol", "Over 4.5 cartellini"], INFO, styled=styled, claude=claude,
                           odds=config.odds, stakes={"1": 10.0})
    assert [list(row) for row in rows] == [list(COLUMNS)] * 3
    one = rows[0]
    assert one["ID"] == "2026-10-25|Inter - Milan|1X2|1" and one["Partita"] == "Inter - Milan"
    assert one["Prob_mercato"] == pytest.approx(base["components"]["mercato"]["1"] / 100)
    assert one["Prob_modello"] == pytest.approx(base["components"]["modello"]["1"] / 100)
    assert one["Prob_stile"] == pytest.approx(styled["components"]["finale"]["1"] / 100)
    assert one["Prob_finale"] == one["Prob_Claude"] == pytest.approx(claude["components"]["finale"]["1"] / 100)
    assert one["Quota_presa"] == 1.9 and one["Puntata"] == 10.0 and one["Tipo"] == "Reale"
    assert one["Quota_chiusura"] is None and one["Esito"] is None
    assert rows[2]["Mercato"] == "Cartellini" and rows[2]["Quota_presa"] is None
    assert rows[2]["Note"] == "motore: Calcolo attuale" and rows[0]["Note"] == "motore: Calcolo attuale"   # fase AN: motore sempre dichiarato
    # Senza Claude né stile la finale è la simulazione di base.
    plain = prediction_rows(base, ["X"], INFO)[0]
    assert plain["Prob_finale"] == pytest.approx(base["components"]["finale"]["X"] / 100)
    assert plain["Prob_stile"] is None and plain["Prob_Claude"] is None
    with pytest.raises(ValueError, match="non previsto"):
        prediction_rows(base, ["Somma gol 3"], INFO)
    # DNB: probabilità condizionata al non rimborso, come la quota.
    dnb = prediction_rows(base, ["Casa DNB"], INFO)[0]
    finale = base["components"]["finale"]
    assert dnb["Mercato"] == "DNB"
    assert dnb["Prob_finale"] == pytest.approx(finale["Casa DNB"] / (100 - finale["Casa DNB push"]), abs=1e-4)


def test_model_version_describes_the_configuration(simulations):
    config = simulations[0]
    assert model_version(config).startswith(config.goal_model)
    styled = replace(config, tactical=engine.TacticalFactors(cards=1.1), tactical_switches=engine.TacticalSwitches(cards=True))
    assert "stile cards" in model_version(styled, claude=True) and model_version(styled, claude=True).endswith("Claude")


def test_session_log_avoids_duplicates_and_migrates_phase_g(tmp_path, simulations):
    _, base, _, _ = simulations
    path = tmp_path / "c.sqlite"
    with sqlite3.connect(path) as connection:
        connection.execute('CREATE TABLE registro_previsioni (id INTEGER PRIMARY KEY AUTOINCREMENT, sessione TEXT, "Data", '
                           '"Partita", "Mercato", "Prob_modello", "Prob_Claude", "Quota", "Confidenza_Claude", "Note")')
        connection.execute('INSERT INTO registro_previsioni VALUES (NULL, "s0", "2026-10-20T19:00:00+00:00", "Napoli - Roma", '
                           '"1", 45.5, 43.0, 2.3, "media", "correzioni di Claude applicate")')
    log = PredictionLog(SourceCache(path, opener=lambda *args: b""))
    migrated = log.rows("s0")
    assert len(migrated) == 1 and migrated[0]["Mercato"] == "1X2" and migrated[0]["Prob_finale"] == 0.43
    assert migrated[0]["Prob_modello"] == 0.455 and "fase G" in migrated[0]["Note"]
    with sqlite3.connect(path) as connection:
        assert connection.execute("SELECT name FROM sqlite_master WHERE name='registro_previsioni'").fetchone() is None

    rows = prediction_rows(base, ["1", "X"], INFO)
    assert log.save(rows, "s1") == (2, 0)
    assert log.save(prediction_rows(base, ["1"], replace(INFO, note="di nuovo")), "s1") == (0, 1)
    stored = log.rows("s1")
    assert len(stored) == 2 and {row["ID"] for row in stored} == {row["ID"] for row in rows}
    assert next(row for row in stored if row["Selezione"] == "1")["Note"] == "di nuovo; motore: Calcolo attuale"


def test_excel_export_has_only_columns_a_to_s(simulations):
    config, base, _, _ = simulations
    rows = prediction_rows(base, ["1", "Over 2.5 gol"], INFO, odds=config.odds)
    sheet = load_workbook(BytesIO(to_excel(rows)))["Registro"]
    assert sheet.max_column == 19 and sheet.max_row == 3
    assert tuple(cell.value for cell in sheet[1]) == COLUMNS
    assert sheet["B2"].value == datetime(2026, 10, 25) and sheet["F3"].value == "Over 2.5"
    assert sheet["H2"].value == rows[0]["Prob_finale"] and sheet["Q2"].value is None and sheet["R2"].value is None
    # Riletto dalla pagina Prestazioni dà le stesse righe.
    registry = read_registry(to_excel(rows))
    assert list(registry.frame["ID"]) == [row["ID"] for row in rows]


# --- Lettura del file dell'utente -----------------------------------------------------------------

def test_empty_template_is_read_without_rows():
    registry = read_registry(TEMPLATE.read_bytes())
    assert registry.sheet == "Registro" and registry.frame.empty


def test_sample_file_with_duplicates_blank_rows_and_empty_columns():
    registry = read_registry(SAMPLE.read_bytes())
    frame = registry.frame
    assert registry.duplicates == 1 and len(frame) == 16  # 18 righe scritte, 1 vuota, 1 doppione
    assert frame["ID"].is_unique
    assert "2026-08-23|Venezia - Lecce|1X2|2" in set(frame["ID"])  # ID ricostruito
    torino = frame[frame["Partita"] == "Torino - Milan"].iloc[0]
    assert torino["Prob_finale"] == pytest.approx(0.62)  # 62 → 0,62
    assert frame.loc[frame["Partita"] == "Atalanta - Sassuolo", "Data"].iloc[0] == date(2026, 8, 23)  # testo gg/mm/aaaa
    assert frame["Prob_Claude"].notna().sum() == 1 and frame["Prob_stile"].notna().sum() == 1
    assert any("percentuale" in warning for warning in registry.warnings)
    assert any("duplicate" in warning for warning in registry.warnings)


def test_missing_columns_are_reported():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Registro"
    workbook.active.append(["ID", "Data", "Partita", "Esito"])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(RegistryError, match="Colonne mancanti"):
        read_registry(buffer.getvalue())
    with pytest.raises(RegistryError, match="non leggibile"):
        read_registry(b"non un file excel")


# --- Esiti e quote di chiusura ------------------------------------------------------------------

def test_settle_uses_engine_event_definitions():
    result = {"home_goals": 2, "away_goals": 1, "home_ht_goals": 0, "away_ht_goals": 1,
              "home_corners": 6, "away_corners": 4, "home_yellows": 3, "away_yellows": 2, "home_reds": 0, "away_reds": 1}
    assert settle("1", result) == 1 and settle("X2", result) == 0
    assert settle("Over 2.5 gol", result) == 1 and settle("Under 2.5 gol", result) == 0
    assert settle("Goal", result) == 1 and settle("Multigol 2-4", result) == 1
    assert settle("2 1T", result) == 1 and settle("Over 9.5 corner", result) == 1
    assert settle("Over 4.5 cartellini", result) == 1 and settle("Multigol Casa 1-2", result) == 1
    # Regola Eurobet: 5 gialli + 1 rosso = 7 punti; ospite 2 + 2 = 4 punti contro i 3 della casa.
    assert settle("Over 6.5 cartellini", result) == 1 and settle("Under 6.5 cartellini", result) == 0
    assert settle("2 cartellini", result) == 1
    assert settle("Over 4.5 cartellini", {**result, "away_reds": np.nan}) is None
    yellows_only = engine.CardRule(red_weight=0)
    assert settle("Over 6.5 cartellini", result, yellows_only) == 0
    assert settle("1 + Over 3.5 gol", result) == 0 and settle("1 + Goal", result) == 1
    assert settle("Over 9.5 corner", {"home_goals": 1, "away_goals": 0, "home_corners": np.nan}) is None
    assert settle("Selezione inesistente", result) is None


def test_split_match_accepts_common_separators():
    assert split_match("Inter - Monza") == ("Inter", "Monza")
    assert split_match("Real Madrid vs Barcellona") == ("Real Madrid", "Barcellona")
    assert split_match("Inter") is None


def test_fill_from_results_completes_only_missing_values():
    registry = read_registry(SAMPLE.read_bytes())
    report = fill_from_results(registry.frame, _results(), today=date(2026, 10, 22))
    rows = report.frame.set_index(["Partita", "Selezione"])

    def cell(match, selection, column):
        return rows.loc[(match, selection), column]

    assert cell("Inter - Monza", "1", "Esito") == 1 and cell("Inter - Monza", "1", "Quota_chiusura") == pytest.approx(1.15)
    note = cell("Inter - Monza", "1", "Note")
    assert "football-data.co.uk" in note and "22/10/2026" in note
    assert cell("Udinese - Como", "Over 2.5 gol", "Esito") == 0
    assert cell("Udinese - Como", "Over 2.5 gol", "Quota_chiusura") == pytest.approx(2.00)
    assert cell("Genoa - Napoli", "Over 9.5 corner", "Esito") == 1
    assert pd.isna(cell("Genoa - Napoli", "Over 9.5 corner", "Quota_chiusura"))
    assert cell("Torino - Milan", "Over 4.5 cartellini", "Esito") == 0  # 1 giallo in tutto
    assert "gialli" not in cell("Torino - Milan", "Over 4.5 cartellini", "Note")
    assert cell("Genoa - Frosinone", "Over 6.5 cartellini", "Esito") == 1  # 5 gialli + 1 rosso = 7 punti
    assert cell("Venezia - Lecce", "2", "Esito") == 1 and cell("Venezia - Lecce", "2", "Quota_chiusura") == pytest.approx(4.54)
    assert cell("Inter - Monza", "1 + Over 1.5 gol", "Esito") == 1
    # Nuovi mercati della fase H2.
    assert cell("Udinese - Como", "1 1T", "Esito") == 1  # 1-0 al 45'
    assert cell("Inter - Monza", "Over 3.5 gol", "Esito") == 1
    assert cell("Inter - Monza", "Over 1.5 gol 1T", "Esito") == 1  # 1-1 al 45'
    assert cell("Genoa - Napoli", "Multigol Ospite 1-2", "Esito") == 1
    assert cell("Torino - Milan", "X2", "Esito") == 1
    # DNB con pareggio (1-1): rimborsata, Esito vuoto e annotato; non viene ricompilata la volta dopo.
    assert pd.isna(cell("Udinese - Como", "Casa DNB", "Esito")) and "rimborsata" in cell("Udinese - Como", "Casa DNB", "Note")
    # Valori già presenti non si toccano.
    assert cell("Atalanta - Sassuolo", "Goal", "Quota_chiusura") == 1.70
    assert cell("Atalanta - Sassuolo", "Goal", "Note") == "esito inserito a mano"
    assert report.not_found == ["Real Madrid - Barcellona"]
    assert cell("Inter - Monza", "3-1/4-1", "Esito") == 1  # finita 4-1
    assert report.outcomes == 13 and report.closings == 3 and report.pushes == 1
    again = fill_from_results(report.frame, _results())
    assert again.pushes == 0 and again.outcomes == 0 and again.changes == []


def test_updated_workbook_keeps_formulas_and_fills_duplicates():
    original = SAMPLE.read_bytes()
    registry = read_registry(original)
    report = fill_from_results(registry.frame, _results(), today=date(2026, 10, 22))
    sheet = load_workbook(BytesIO(updated_workbook(original, registry.sheet, report.changes)))["Registro"]
    assert sheet["R2"].value == 1 and sheet["R3"].value == 1  # stessa partita incollata due volte
    assert sheet["Q2"].value == pytest.approx(1.15) and sheet["S3"].value.endswith("22/10/2026")
    assert sheet["T2"].value.startswith("=IF(") and sheet["Y3001"].value.startswith("=IF(")
    assert sheet["R6"].value == 1 and sheet["S6"].value == "esito inserito a mano"
    assert sheet["R10"].value is None  # Champions League: non trovata
    assert load_workbook(BytesIO(original))["Registro"]["R2"].value is None  # l'originale non cambia


def test_collect_results_downloads_only_needed_leagues_and_seasons():
    registry = read_registry(SAMPLE.read_bytes())
    calls = []

    def fetch(league, season):
        calls.append((league, season))
        return _results() if league == "I1" else pd.DataFrame()

    results, warnings = collect_results(registry.frame, fetch, COMPETITIONS)
    # Una riga senza competizione: si cercano tutti i campionati coperti, stagione 2026-27.
    assert {season for _, season in calls} == {2026} and len(calls) == len(COMPETITIONS)
    assert not results.empty and any("Champions League" in warning for warning in warnings)
    only_serie_a = registry.frame[registry.frame["Competizione"] == "Serie A"]
    calls.clear()
    collect_results(only_serie_a, fetch, COMPETITIONS)
    assert calls == [("I1", 2026)]


# --- Misure ----------------------------------------------------------------------------------------

def _synthetic(count: int = 240, seed: int = 3) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    market = rng.uniform(0.3, 0.8, count)
    outcome = (rng.uniform(size=count) < market).astype(float)
    taken = rng.uniform(1.5, 2.5, count)
    return pd.DataFrame({
        "ID": [f"r{i}" for i in range(count)], "Data": [date(2026, 8, 1) + pd.Timedelta(days=int(i // 4)) for i in range(count)],
        "Ora": "20:45", "Competizione": "Serie A", "Partita": "A - B",
        "Mercato": np.where(np.arange(count) < 150, "1X2", "Corner"), "Selezione": "1",
        "Prob_finale": market, "Prob_mercato": market, "Prob_modello": np.clip(market + rng.normal(0, 0.15, count), 0.01, 0.99),
        "Prob_stile": np.nan, "Prob_Claude": np.nan, "Quota_presa": taken, "Puntata": 10.0,
        "Tipo": np.where(np.arange(count) % 2 == 0, "Reale", "A carta"), "Versione_modello": "x",
        "Quota_chiusura": np.where(np.arange(count) < 150, taken / 1.03, taken * 1.02), "Esito": outcome, "Note": "",
    })


def test_scores_and_calibration():
    frame = _synthetic()
    assert brier([1.0, 0.0], [1, 0]) == 0 and log_loss([0.5], [1]) == pytest.approx(np.log(2))
    table = calibration_table(frame)
    assert set(table["Mercato"]) == {"1X2", "Corner"} and table["Selezioni"].sum() == len(frame)
    assert set(table["Fascia"]) == {"0%–50%", "50%–60%", "60%–70%", "70%–80%"}
    monthly = scores_over_time(frame, "M")
    assert monthly["Selezioni"].sum() == len(frame)
    assert monthly["Brier cumulato"].iloc[-1] == pytest.approx(brier(frame["Prob_finale"], frame["Esito"]))
    assert len(scores_over_time(frame, "W")) > len(monthly)


def test_clv_roi_and_verdicts():
    frame = _synthetic()
    clv = clv_by_market(frame).set_index("Mercato")
    assert clv.loc["1X2", "CLV medio"] == pytest.approx(0.03) and clv.loc["Corner", "CLV medio"] == pytest.approx(-1 / 51)
    summary, timeline = bankroll(frame)
    summary = summary.set_index("Tipo")
    profit = 10.0 * (frame["Quota_presa"] * frame["Esito"] - 1)
    real = frame["Tipo"] == "Reale"
    assert summary.loc["Reale", "Profitto"] == pytest.approx(profit[real].sum())
    assert summary.loc["A carta", "ROI"] == pytest.approx(profit[~real].sum() / (10.0 * (~real).sum()))
    assert timeline.groupby("Tipo")["Saldo"].last()["Reale"] == pytest.approx(profit[real].sum())
    verdicts = edge_verdicts(frame).set_index("Mercato")
    assert verdicts.loc["1X2", "Giudizio"] == "VANTAGGIO DIMOSTRATO"
    assert verdicts.loc["Corner", "Giudizio"] == "Non ancora dimostrato (servono 100, ne hai 90)"
    assert list(verdicts.index[:len(MARKETS)]) == list(MARKETS)
    many_negative = _synthetic(400)
    many_negative["Mercato"] = "Corner"
    assert edge_verdicts(many_negative).set_index("Mercato").loc["Corner", "Giudizio"] == "Nessun vantaggio (CLV non positivo)"


def test_component_comparison_detects_a_worse_model():
    frame = _synthetic(600)
    table = component_comparison(frame).set_index("Confronto")
    assert table.loc["Modello vs Mercato", "Giudizio"] == "peggiora"
    assert table.loc["Modello", "Brier"] > table.loc["Mercato", "Brier"]
    assert table.loc["Stile", "Selezioni"] == 0 and "Stile vs Mercato" not in table.index
    better = frame.assign(Prob_Claude=frame["Esito"] * 0.6 + 0.2)
    assert component_comparison(better).set_index("Confronto").loc["Claude vs Mercato", "Giudizio"] == "migliora davvero"


def test_measures_on_the_sample_file_after_filling():
    registry = read_registry(SAMPLE.read_bytes())
    frame = fill_from_results(registry.frame, _results()).frame
    verdicts = edge_verdicts(frame).set_index("Mercato")
    assert verdicts.loc["1X2", "Selezioni con esito"] == 2  # il doppione conta una volta
    summary, _ = bankroll(frame)
    assert set(summary["Tipo"]) == {"Reale", "A carta"}
    assert not component_comparison(frame).empty and not calibration_table(frame).empty


def test_exact_score_group_rows_use_the_score_distribution(simulations):
    config, base, _, claude = simulations
    rows = prediction_rows(base, ["2-1/1-0/2-0"], INFO, claude=claude, odds={"2-1/1-0/2-0": 3.2})
    row = rows[0]
    assert row["Mercato"] == "Risultati esatti" and row["Selezione"] == "2-1/1-0/2-0"
    expected = sum(claude["exact_scores"]["finale"].get(score, 0) for score in ("1-0", "2-0", "2-1")) / 100
    assert row["Prob_finale"] == pytest.approx(expected, abs=1e-4)
    model = sum(base["exact_scores"]["modello"].get(score, 0) for score in ("1-0", "2-0", "2-1")) / 100
    assert row["Prob_modello"] == pytest.approx(model, abs=1e-4) and row["Prob_mercato"] is not None
    assert row["Quota_presa"] == 3.2
    assert settle("1-0/2-0/2-1", {"home_goals": 2, "away_goals": 0}) == 1
    assert settle("1-0/2-0/2-1", {"home_goals": 0, "away_goals": 0}) == 0
    assert settle("1-0/2-0", {"home_goals": np.nan, "away_goals": 0}) is None
